"""The Portfolio download: exactly the rows and columns on screen, as Excel or PDF.

Asked for as *"if I have just selected sattu… download only sattu"*, *"text which are numbers stored
as numbers or %"*, *"the same format as what we see on the screen"*, *"no additional commentary in
the excel. just data"* and *"give both pdf and excel option"*.

**The browser chooses the ROWS; this module computes every NUMBER.** The screen sends the ids its
own `visible()` returned — so the category card, group tab, custom filters, search and grain are
all honoured with no second copy of that logic — and the figures come from `_dashboard`, the same
builder the screen reads. Nothing numeric travels from the client.

Measured on the workbook this replaces: percentages, weights and ratings were TEXT (`'30.0%'`,
`'1,453.0 kg'`, `'4.0 (388)'`), row 1 was a ~600-character sentence, size rows carried prose in the
Brand column, and every cell was `General`. Here every figure is a number with a number format
that makes it READ like the screen (₹ with lakh grouping, `0.0%`, `kg`), and a missing value is a
blank cell — the screen's "a dash, never a zero", since 0% TACOS would read as perfectly efficient.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from app.portfolio import columns as column_vocab

#: How each screen column is stored and shown. Every id in `columns.COLUMNS` except `product` must
#: be here — a test asserts it, so a new screen column cannot silently vanish from the download.
COLUMN_KINDS: dict[str, str] = {
    "verdict": "text", "sales": "money", "ad_spend": "money",
    "refunds_pct": "pct", "fees_pct": "pct", "tacos": "pct", "net_pct": "pct_signed",
    "acos": "acos", "units_ordered": "int", "units": "int",
    "weight_ordered_kg": "kg", "weight_kg": "kg", "returns_pct": "pct",
    "rating": "rating", "decision": "text",
}
LEAD = [("product", "Product", "text"), ("brand", "Brand", "text"),
        ("size", "Size", "text"), ("asin", "ASIN", "text")]

# Indian digit grouping (12,34,567) like the screen's `toLocaleString("en-IN")`. Excel has no
# locale switch for it, so the two thresholds a lakh and a crore are written into the format.
def _indian_format(prefix: str = "", suffix: str = "") -> str:
    return (f'[>=10000000]{prefix}##\\,##\\,##\\,##0{suffix};'
            f'[>=100000]{prefix}##\\,##\\,##0{suffix};'
            f'{prefix}##,##0{suffix}')


EXCEL_FORMATS = {
    "money": _indian_format(prefix='"₹"'),
    "int": _indian_format(),
    "kg": _indian_format(suffix='.0" kg"'),
    "pct": "0.0%",
    "pct_signed": "+0.0%;-0.0%;0.0%",
    "acos": "0%",
    "rating": "0.0",
    "reviews": _indian_format(),
    "text": "General",
}

FONTS = Path(__file__).resolve().parent.parent / "fonts"

#: Hover notes on headings: the caveats that change how a figure is read.
HEADER_NOTES = {
    "net_pct": "PRE-COGS: 100% − Refunds % − Amazon fees % − TACOS, before what it costs to make "
               "the product.",
    "fees_pct": "Every Amazon fee except advertising, after waivers, EXCLUDING the 18% GST "
                "(claimed as ITC).",
}


@dataclass
class Column:
    id: str
    header: str
    kind: str


@dataclass
class Row:
    kind: str                 # "parent" | "flavour" | "size" | "sku"
    level: int                # outline level: 0 parent/sku, 1 size or flavour, 2 size in a flavour
    hidden: bool              # collapsed on screen
    values: dict = field(default_factory=dict)


@dataclass
class Table:
    columns: list[Column]
    rows: list[Row]
    total_label: str
    totals: dict


# ── values ──────────────────────────────────────────────────────────────────────────────────────

def _ratio(part, whole):
    """The page's `ratio`: no denominator is None, never 0."""
    return part / whole if whole else None


def _num(v) -> float:
    return float(v or 0)


def _size_label(row: Mapping) -> str:
    from app.shipment.logic import weight_label

    weight = float(row.get("weight") or 0)
    return weight_label(weight) if weight else ""


def _metric(col_id: str, row: Mapping):
    """One figure for one row, typed. Mirrors the page's `row` / `detail` renderers."""
    if col_id == "acos":
        if row.get("acos_infinite"):
            return "no sales"
        return row.get("acos")
    if col_id == "returns_pct":
        return row.get("returns_pct") or None          # the screen dashes a zero here
    if col_id in ("units", "units_ordered"):
        return int(row.get(col_id) or 0)
    return row.get(col_id)


def _values(row: Mapping, kind: str, ids: Sequence[str]) -> dict:
    out = {
        "product": (row.get("flavour") if kind == "flavour" else row.get("product")) or "",
        "brand": row.get("brand") or "",
        "size": _size_label(row) if kind in ("size", "sku") else "",
        "asin": (row.get("asin") if kind in ("size", "sku") else row.get("parent_asin")) or "",
    }
    for col_id in ids:
        if col_id == "verdict":
            out[col_id] = row.get("verdict") if kind in ("parent", "sku") else None
        elif col_id == "rating":
            # A size has no rating of its own: Amazon pools reviews per family.
            rated = kind in ("parent", "sku") and row.get("rating") is not None
            out["rating"] = row.get("rating") if rated else None
            out["reviews"] = int(row.get("rating_count") or 0) if rated else None
        elif col_id == "decision":
            d = row.get("decision") if kind in ("parent", "sku") else None
            out[col_id] = d.upper() if d else None
        else:
            out[col_id] = _metric(col_id, row)
    return out


def totals(rows: Sequence[Mapping]) -> dict:
    """The page's `computeTotals`, rule for rule — and pinned to it by a test that runs the page's
    own JavaScript on the same rows. Money and units SUM; every percentage is recomputed from the
    sums, never averaged; weight counts only rows that HAVE one; rating is weighted by reviews and
    deduplicated per family (Amazon pools reviews, so summing per size claimed 16,789 reviews where
    4,382 exist)."""
    s = lambda key: sum(_num(r.get(key)) for r in rows)  # noqa: E731
    sales, spend, net = s("sales"), s("ad_spend"), s("net")
    fees, refunded_sales = s("fees_total"), s("refunded")
    ads_cost, attributed = s("ads_cost"), s("ad_attributed_sales")
    units, ordered, refunded = s("units"), s("units_ordered"), s("units_refunded")
    weighed = [r for r in rows if r.get("weight_kg") is not None]
    weighed_ordered = [r for r in rows if r.get("weight_ordered_kg") is not None]
    seen, rated = set(), []
    for r in rows:
        if r.get("rating") is None:
            continue
        family = r.get("parent_asin") or r.get("asin")
        if family in seen:
            continue
        seen.add(family)
        rated.append(r)
    reviews = sum(_num(r.get("rating_count")) for r in rated)
    rating = (sum(r["rating"] * _num(r.get("rating_count")) for r in rated) / reviews
              if reviews else None)
    decided = sum(1 for r in rows if r.get("decision"))
    infinite = ads_cost > 0 and attributed == 0
    return {
        "sales": sales, "ad_spend": spend,
        "refunds_pct": _ratio(refunded_sales, sales), "fees_pct": _ratio(fees, sales),
        "tacos": _ratio(spend, sales), "net_pct": _ratio(net, sales),
        "acos": "no sales" if infinite else (_ratio(ads_cost, attributed) if ads_cost else None),
        "units_ordered": int(ordered), "units": int(units),
        "weight_ordered_kg": sum(r["weight_ordered_kg"] for r in weighed_ordered)
                             if weighed_ordered else None,
        "weight_kg": sum(r["weight_kg"] for r in weighed) if weighed else None,
        "returns_pct": _ratio(refunded, ordered),
        "rating": rating, "reviews": int(reviews) if rating is not None else None,
        "decision": f"{decided} decided" if decided else None,
        "verdict": None,
    }


# ── the table ───────────────────────────────────────────────────────────────────────────────────

def default_columns() -> list[str]:
    return [c["id"] for c in column_vocab.COLUMNS if c["id"] != "product"]


def build_table(data: Mapping, *, view: str = "products", ids: Sequence[str] | None = None,
                open_ids: Iterable[str] = (), columns: Sequence[str] | None = None) -> Table:
    """The rows the screen shows, in its order, with its columns.

    `ids` None means every row (the unfiltered GET). Unknown ids and unknown column ids are dropped
    rather than refused: a stale tab must still get a file, and nothing unknown can reach a cell.
    """
    wanted = [c for c in (columns if columns is not None else default_columns())
              if c in COLUMN_KINDS]
    wanted = list(dict.fromkeys(wanted))
    labels = {c["id"]: c["label"] for c in column_vocab.COLUMNS}
    cols = [Column(i, h, k) for i, h, k in LEAD]
    for col_id in wanted:
        cols.append(Column(col_id, labels.get(col_id, col_id), COLUMN_KINDS[col_id]))
        if col_id == "rating":
            cols.append(Column("reviews", "Reviews", "reviews"))

    skus = view == "skus"
    source = list(data.get("skus" if skus else "parents") or [])
    key = (lambda r: r.get("asin")) if skus else (lambda r: r.get("parent_asin"))
    if ids is None:
        chosen = source
    else:
        by_id = {key(r): r for r in source}
        chosen = [by_id[i] for i in dict.fromkeys(ids) if i in by_id]
    opened = set(open_ids or ())

    rows: list[Row] = []
    for r in chosen:
        if skus:
            rows.append(Row("sku", 0, False, _values(r, "sku", wanted)))
            continue
        rows.append(Row("parent", 0, False, _values(r, "parent", wanted)))
        collapsed = r.get("parent_asin") not in opened
        groups = r.get("flavour_groups") or []
        if groups:
            for g in groups:
                rows.append(Row("flavour", 1, collapsed,
                                _values({**g, "brand": r.get("brand")}, "flavour", wanted)))
                for size in g.get("sizes") or []:
                    rows.append(Row("size", 2, collapsed,
                                    _values({**size, "brand": size.get("brand") or r.get("brand")},
                                            "size", wanted)))
        else:
            for size in r.get("sizes") or []:
                rows.append(Row("size", 1, collapsed,
                                _values({**size, "brand": size.get("brand") or r.get("brand")},
                                        "size", wanted)))
    noun = "pack size(s)" if skus else "product(s)"
    return Table(cols, rows, f"Total — {len(chosen):,} {noun}", totals(chosen))


# ── Excel ───────────────────────────────────────────────────────────────────────────────────────

def build_xlsx(table: Table, sheet: str = "Portfolio") -> io.BytesIO:
    """Row 1 totals, row 2 headings with filter buttons, row 3+ data. No title, no commentary."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    size_font = Font(name="Arial", size=9, color="595959")
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    total_fill = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="BFBFBF")
    edge = Border(bottom=thin)

    def put(r, c, value, col, font, fill=None, indent=0):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = font
        if fill:
            cell.fill = fill
        if col.kind != "text":
            cell.number_format = EXCEL_FORMATS[col.kind]
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(vertical="center", indent=indent)
        cell.border = edge
        return cell

    # Row 1: the totals, as the screen shows them above the rows.
    for c, col in enumerate(table.columns, 1):
        value = table.total_label if c == 1 else table.totals.get(col.id)
        put(1, c, value, col, bold, total_fill)
    ws.cell(row=1, column=1).comment = Comment(
        "Calculated by the app for the rows below, exactly as the Portfolio screen totals them: "
        "money and units are summed, percentages are recomputed from those sums (never averaged), "
        "and rating is weighted by reviews.", "Amazon Tracker")

    # Row 2: headings. The two caveats that change how a figure is READ travel as hover notes on
    # their headings — not as rows of commentary, which the owner asked to be rid of, and not
    # dropped, because a forwarded file showing "+8.8% net" otherwise reads as profit.
    for c, col in enumerate(table.columns, 1):
        cell = ws.cell(row=2, column=c, value=col.header)
        if col.id in HEADER_NOTES:
            cell.comment = Comment(HEADER_NOTES[col.id], "Amazon Tracker")
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="808080"))
    ws.row_dimensions[2].height = 30

    widths = [len(str(col.header)) for col in table.columns]
    for i, row in enumerate(table.rows, 3):
        font = bold if row.kind == "parent" else (size_font if row.kind == "size" else base)
        for c, col in enumerate(table.columns, 1):
            value = row.values.get(col.id)
            put(i, c, value, col, font, indent=row.level if c == 1 else 0)
            if col.kind == "text" and value:
                widths[c - 1] = max(widths[c - 1], len(str(value)) + 2 * (row.level if c == 1 else 0))
        if row.level:
            ws.row_dimensions[i].outlineLevel = row.level
            ws.row_dimensions[i].hidden = row.hidden
    ws.sheet_properties.outlinePr.summaryBelow = False

    for c, col in enumerate(table.columns, 1):
        if col.kind == "text":
            w = min(max(widths[c - 1], len(col.header)) + 2, 42)
        else:
            w = max(11, min(len(col.header) + 2, 16))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 28)

    last = get_column_letter(len(table.columns))
    ws.auto_filter.ref = f"A2:{last}{max(2, len(table.rows) + 2)}"
    ws.freeze_panes = "E3"
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ── PDF ─────────────────────────────────────────────────────────────────────────────────────────

def indian(n: float) -> str:
    """12,34,567 grouping, as the screen's `toLocaleString("en-IN")`."""
    neg, s = n < 0, str(int(round(abs(n))))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts + [tail])
    return ("-" if neg else "") + s


def display(value, kind: str) -> str:
    """A cell as the SCREEN prints it. Blank for a missing value."""
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value
    if kind == "money":
        return ("-" if value < 0 else "") + "₹" + indian(abs(value))
    if kind in ("int", "reviews"):
        return indian(value)
    if kind == "pct":
        return f"{value * 100:.1f}%"
    if kind == "pct_signed":
        return f"{'+' if value > 0 else ''}{value * 100:.1f}%"
    if kind == "acos":
        return f"{value * 100:.0f}%"
    if kind == "kg":
        tenth = round(abs(value), 1)
        whole, frac = int(tenth), round(tenth - int(tenth), 1)
        return ("-" if value < 0 else "") + indian(whole) + (f".{int(round(frac * 10))}" if frac else "") + " kg"
    if kind == "rating":
        return f"{value:.1f}"
    return str(value)


def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVuCond" not in pdfmetrics.getRegisteredFontNames():
        # reportlab's base fonts have no ₹, and the production box has no TrueType fonts at all.
        pdfmetrics.registerFont(TTFont("DejaVuCond", str(FONTS / "DejaVuSansCondensed.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVuCond-Bold", str(FONTS / "DejaVuSansCondensed-Bold.ttf")))
    return "DejaVuCond", "DejaVuCond-Bold"


def build_pdf(table: Table, title: str) -> io.BytesIO:
    """Landscape A4: headings, then totals, then the rows the screen shows (collapsed sizes are
    left out, as on screen). Every cell a Paragraph — a bare string overflows its gridline."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as PTable, TableStyle

    regular, bold = _fonts()

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    shown = [r for r in table.rows if not r.hidden]
    grid = [[col.header for col in table.columns]]
    grid.append([table.total_label if i == 0 else display(table.totals.get(c.id), c.kind)
                 for i, c in enumerate(table.columns)])
    for r in shown:
        grid.append([display(r.values.get(c.id), c.kind) for c in table.columns])

    page_w = landscape(A4)[0] - 20 * mm
    # Text that may wrap (product and brand names, the total's label) gives up width first. Every
    # other column is NEVER narrower than its widest value, and a heading wraps only between words:
    # proportional shrinking printed "₹8,48,22 / 5" and "ACO / S" on the first real download.
    flexible = {"product", "brand"}

    def widths_at(size: float):
        pad = 6
        out, floors = [], []
        for i, col in enumerate(table.columns):
            body = [str(row[i]) for row in grid[1:]] or [""]
            widest = max(stringWidth(v, bold, size) for v in body)
            word = max(stringWidth(w, bold, size) for w in str(col.header).split())
            if col.id in flexible:
                out.append(min(widest, 150) + pad)
                floors.append(max(word, 55) + pad)
            else:
                need = max(widest, word) + pad
                out.append(need)
                floors.append(need)
        return out, floors

    for size in (7, 6.5, 6, 5.5):
        natural, floors = widths_at(size)
        excess = sum(natural) - page_w
        if excess <= 0:
            widths = natural
            break
        give = {i: natural[i] - floors[i] for i, c in enumerate(table.columns) if c.id in flexible}
        room = sum(give.values())
        if room >= excess:
            widths = [w - (give.get(i, 0) / room * excess if room else 0) for i, w in enumerate(natural)]
            break
    else:
        widths = [w * page_w / sum(natural) for w in natural]
    left = ParagraphStyle("l", fontName=regular, fontSize=size, leading=size + 2)
    right = ParagraphStyle("r", parent=left, alignment=2)
    head = ParagraphStyle("h", parent=left, fontName=bold, alignment=1)

    cells = []
    for ri, row in enumerate(grid):
        line = []
        for ci, (value, col) in enumerate(zip(row, table.columns)):
            if ri == 0:
                line.append(Paragraph(esc(str(value)), head))
                continue
            style = right if col.kind != "text" else left
            if ri == 1:
                style = ParagraphStyle("t", parent=style, fontName=bold)
            elif ci == 0 and shown[ri - 2].level:
                style = ParagraphStyle("i", parent=left, leftIndent=6 * shown[ri - 2].level,
                                       textColor=colors.HexColor("#555555"))
            line.append(Paragraph(esc(str(value)), style))
        cells.append(line)

    t = PTable(cells, colWidths=widths, repeatRows=1)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EDEDED")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
    ]
    for ri, r in enumerate(shown, 2):
        if r.kind == "parent":
            style.append(("FONTNAME", (0, ri), (-1, ri), bold))
    t.setStyle(TableStyle(style))

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=title)
    heading = ParagraphStyle("title", fontName=bold, fontSize=10, leading=13)
    doc.build([Paragraph(esc(title), heading), Spacer(1, 3 * mm), t])
    out.seek(0)
    return out


def filename(window: Sequence[str] | None, label: str, ext: str) -> str:
    """`portfolio-<start>_<end>[-<filter>].<ext>` — the filter slugged so it is a safe filename."""
    stem = "portfolio"
    if window:
        stem += f"-{window[0]}_{window[1]}"
    slug = re.sub(r"[^a-z0-9]+", "-", (label or "").lower()).strip("-")[:60]
    return f"{stem}{'-' + slug if slug else ''}.{ext}"
