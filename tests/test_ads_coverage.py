"""Ads → SKU coverage: every SKU of an advertised product is in its SP ad group (24 Oct 2026).

Asked for as *"check that all the sku's of a parent has been added to that particular ad group where
it should — whether it is active or inactive is a different thing but it should be added. in SP…
maybe we do the list in fba and flex but forgot to add it in the ad group where the easy ship sku was
added."* Measured before building: 8 enabled ad groups missed 77 SKUs of ASINs they already advertise.
"""
from datetime import datetime, timedelta
from pathlib import Path

import pytest

from app.ads import coverage, listings, logic, repository, spapi_ads

ROOT = Path(__file__).resolve().parent.parent
PAGE = (ROOT / "templates" / "ads_expansion.html").read_text(encoding="utf-8")
OPTIMISATION = (ROOT / "templates" / "ads.html").read_text(encoding="utf-8")

CAMPAIGNS = {"c1": {"name": "MF_SP_keywords", "state": "ENABLED"},
             "c2": {"name": "m19 autopilot chana", "state": "ENABLED"},
             "c3": {"name": "old", "state": "PAUSED"}}
GROUPS = {"g1": {"name": "chana sattu", "state": "ENABLED"},
          "g2": {"name": "roasted chana all", "state": "ENABLED"},
          "g3": {"name": "paused group", "state": "PAUSED"},
          "g4": {"name": "m19", "state": "ENABLED"},
          "g5": {"name": "old", "state": "ENABLED"}}


def _listing(sku, asin, status="Active"):
    return {"seller_sku": sku, "asin": asin, "status": status, "fulfillment_channel": "DEFAULT",
            "item_name": ""}


LISTINGS = [
    # Chana Sattu: two sizes, three channels each.
    _listing("0.5kg cs 1", "A500"), _listing("0.5kg cs 1 FBA", "A500"),
    _listing("0.5kg cs 1 flex", "A500", "Inactive"),
    _listing("1kg cs 1", "A1KG"), _listing("1kg cs 1 FBA", "A1KG"),
    _listing("1kg cs 1 flex", "A1KG", "Incomplete"),
    # Retired 2 kg: marked N in the MRP sheet.
    _listing("2kg cs 1", "A2KG"),
    # Roasted chana: two flavours under ONE Amazon parent.
    _listing("pp 250", "PERI"), _listing("pp 250 FBA", "PERI"),
    _listing("np 250", "NIMBU"), _listing("np 250 FBA", "NIMBU"),
]
CATALOGUE = {
    "A500": {"name": "Chana Sattu", "weight": 0.5, "active": True},
    "A1KG": {"name": "Chana Sattu", "weight": 1, "active": True},
    "A2KG": {"name": "Chana Sattu", "weight": 2, "active": False},
    "PERI": {"name": "Peri Peri Roasted Chana", "weight": 0.25, "active": True},
    "NIMBU": {"name": "Nimbu Pudina Roasted Chana", "weight": 0.25, "active": True},
}
# The Portfolio grouping: one Chana Sattu row; the flavours split into their own rows.
ROW = {"A500": "P-SATTU", "A1KG": "P-SATTU", "A2KG": "P-SATTU", "PERI": "PERI", "NIMBU": "NIMBU"}


def _ad(ad_id, group, sku, asin, state="ENABLED", campaign="c1", product="sp"):
    return {"entity_id": ad_id, "parent_id": group, "campaign_id": campaign, "name": sku,
            "asin": asin, "state": state, "ad_product": product}


def _compute(ads, ignores=None, listings_=LISTINGS, catalogue=CATALOGUE):
    return coverage.compute(ads, CAMPAIGNS, GROUPS, listings_, lambda a: ROW.get(a, a),
                            catalogue, ignores or set())


def _missing(result, group="g1"):
    return {m["sku"]: m for c in result["campaigns"] for g in c["groups"] if g["ad_group_id"] == group
            for p in g["products"] for m in p["missing"]}


# ─── The rule ────────────────────────────────────────────────────────────────


def test_a_forgotten_easy_ship_sku_of_an_advertised_asin_is_named():
    result = _compute([_ad("1", "g1", "0.5kg cs 1 FBA", "A500"), _ad("2", "g1", "0.5kg cs 1 flex", "A500")])
    missing = _missing(result)
    assert missing["0.5kg cs 1"]["kind"] == coverage.KIND_SKU
    assert missing["0.5kg cs 1"]["channel"] == logic.CHANNEL_EASY_SHIP
    assert missing["0.5kg cs 1"]["size"] == "500g"


def test_a_size_the_group_does_not_hold_is_a_SIZE_gap():
    missing = _missing(_compute([_ad("1", "g1", "0.5kg cs 1", "A500")]))
    assert missing["1kg cs 1"]["kind"] == coverage.KIND_SIZE
    assert missing["1kg cs 1 FBA"]["kind"] == coverage.KIND_SIZE
    assert missing["0.5kg cs 1 FBA"]["kind"] == coverage.KIND_SKU


def test_an_inactive_listing_is_still_expected_but_an_incomplete_one_is_not():
    result = _compute([_ad("1", "g1", "0.5kg cs 1", "A500"), _ad("2", "g1", "1kg cs 1", "A1KG")])
    missing = _missing(result)
    assert missing["0.5kg cs 1 flex"]["listing_status"] == "Inactive"
    assert "1kg cs 1 flex" not in missing
    assert result["totals"]["incomplete_skipped"] == 1


def test_a_size_retired_in_the_MRP_sheet_is_never_expected():
    assert "2kg cs 1" not in _missing(_compute([_ad("1", "g1", "0.5kg cs 1", "A500")]))


def test_a_missing_active_flag_is_not_a_retirement():
    catalogue = {**CATALOGUE, "A2KG": {"name": "Chana Sattu", "weight": 2}}
    assert "2kg cs 1" in _missing(_compute([_ad("1", "g1", "0.5kg cs 1", "A500")], catalogue=catalogue))


def test_a_flavours_own_group_expects_only_that_flavour():
    missing = _missing(_compute([_ad("1", "g2", "pp 250", "PERI")]), "g2")
    assert set(missing) == {"pp 250 FBA"}


def test_a_group_holding_every_flavour_expects_every_flavours_skus():
    missing = _missing(_compute([_ad("1", "g2", "pp 250", "PERI"), _ad("2", "g2", "np 250", "NIMBU")]), "g2")
    assert set(missing) == {"pp 250 FBA", "np 250 FBA"}


def test_a_paused_sku_counts_as_present_and_an_archived_one_does_not():
    ads = [_ad("1", "g2", "pp 250", "PERI"), _ad("2", "g2", "pp 250 FBA", "PERI", state="PAUSED")]
    assert _missing(_compute(ads), "g2") == {}
    ads[1]["state"] = "ARCHIVED"
    assert set(_missing(_compute(ads), "g2")) == {"pp 250 FBA"}


def test_paused_groups_paused_campaigns_managed_campaigns_and_SB_are_not_checked():
    ads = [_ad("1", "g3", "pp 250", "PERI"), _ad("2", "g4", "pp 250", "PERI", campaign="c2"),
           _ad("3", "g5", "pp 250", "PERI", campaign="c3"),
           _ad("4", "g2", "pp 250", "PERI", product="sb")]
    result = _compute(ads)
    assert result["campaigns"] == [] and result["totals"]["groups_checked"] == 0
    assert result["skipped"] == {"campaign_not_enabled": 1, "managed_elsewhere": 1, "ad_group_paused": 1}


def test_an_ignored_product_or_group_is_flagged_and_left_out_of_the_counts():
    ads = [_ad("1", "g1", "0.5kg cs 1", "A500"), _ad("2", "g2", "pp 250", "PERI")]
    full = _compute(ads)["totals"]["missing"]
    result = _compute(ads, {("g2", "PERI")})
    assert result["totals"]["missing"] == full - 1 and result["totals"]["ignored"] == 1
    g2 = next(g for c in result["campaigns"] for g in c["groups"] if g["ad_group_id"] == "g2")
    assert g2["products"][0]["ignored"] and g2["missing"] == 0
    assert ("g2", "pp 250 FBA") not in coverage.gaps_by_key(result)
    whole = _compute(ads, {("g1", coverage.WHOLE_GROUP)})
    assert all(k[0] != "g1" for k in coverage.gaps_by_key(whole))


def test_the_asin_comes_from_the_listings_when_the_cache_has_none():
    missing = _missing(_compute([_ad("1", "g1", "0.5kg cs 1", "")]))
    assert missing["0.5kg cs 1 FBA"]["kind"] == coverage.KIND_SKU


def test_the_counts_are_the_rows_on_screen():
    result = _compute([_ad("1", "g1", "0.5kg cs 1", "A500"), _ad("2", "g2", "pp 250", "PERI")])
    shown = sum(len(p["missing"]) for c in result["campaigns"] for g in c["groups"] for p in g["products"])
    assert shown == result["totals"]["missing"] == len(coverage.gaps_by_key(result))
    assert sum(result["totals"]["by_kind"].values()) == result["totals"]["missing"]
    assert result["totals"]["groups_with_gaps"] == 2


def test_the_provisional_id_fits_the_ledger_and_is_stable():
    pid = coverage.provisional_id("406048686897409", "HF_PanPho_0.2kg_flex")
    assert pid.startswith("new-") and len(pid) <= 32
    assert pid == coverage.provisional_id("406048686897409", "HF_PanPho_0.2kg_flex")
    assert pid != coverage.provisional_id("406048686897409", "HF_PanPho_0.5kg_flex")


# ─── The listings ────────────────────────────────────────────────────────────


def test_the_listings_report_is_parsed_through_its_byte_order_mark():
    rows = [{"\ufeffitem-name": "Chana Sattu", "seller-sku": "0.5kg cs 1", "asin1": "A500",
             "status": "Active", "fulfillment-channel": "DEFAULT"},
            {"\ufeffitem-name": "dup", "seller-sku": "0.5kg cs 1", "asin1": "OTHER"},
            {"\ufeffitem-name": "blank", "seller-sku": "", "asin1": "A1KG"}]
    assert listings.parse_rows(rows) == [{"seller_sku": "0.5kg cs 1", "asin": "A500", "status": "Active",
                                          "fulfillment_channel": "DEFAULT", "item_name": "Chana Sattu"}]


async def test_an_empty_read_never_wipes_the_stored_listings(db):
    assert await listings.replace_listings(db, [_listing("a", "A500")]) == 1
    assert await listings.replace_listings(db, []) == 0
    assert [r["seller_sku"] for r in await listings.load_listings(db)] == ["a"]


async def test_a_failed_read_is_recorded_and_the_previous_list_stands(db_schema, monkeypatch):
    from app.database import async_session

    async def good(client, **kw):
        return [_listing("a", "A500")]
    monkeypatch.setattr(listings, "fetch", good)
    first = await listings.refresh(async_session, None)
    assert first["rows"] == 1 and first["ok_at"] and not first["error"]
    # An older good read, so a failure that overwrote the time would be visible.
    import json

    from app.models import PortfolioSettings
    from sqlalchemy import select
    async with async_session() as db:
        row = (await db.execute(select(PortfolioSettings).where(
            PortfolioSettings.name == listings.STATUS_SETTING))).scalar_one()
        row.value_json = json.dumps({**first, "ok_at": "2026-01-01T00:00:00"})
        await db.commit()

    async def bad(client, **kw):
        raise listings.ListingsFailed("quota")
    monkeypatch.setattr(listings, "fetch", bad)
    second = await listings.refresh(async_session, None)
    assert "quota" in second["error"] and second["ok_at"] == "2026-01-01T00:00:00"
    async with async_session() as db:
        assert len(await listings.load_listings(db)) == 1


async def test_the_refresh_reads_the_listings_only_when_stale(db_schema, monkeypatch):
    from app.database import async_session
    calls = []

    async def good(client, **kw):
        calls.append(1)
        return [_listing("a", "A500")]
    monkeypatch.setattr(listings, "fetch", good)
    await listings.refresh_if_stale(async_session, None)
    await listings.refresh_if_stale(async_session, None)
    assert len(calls) == 1
    await listings.refresh_if_stale(async_session, None,
                                    now=datetime.utcnow() + timedelta(hours=listings.STALE_AFTER_HOURS + 1))
    assert len(calls) == 2


def test_the_ads_refresh_reads_the_listings():
    source = (ROOT / "app" / "ads" / "refresh.py").read_text(encoding="utf-8")
    assert "await listings.refresh_if_stale(db_factory, client" in source


# ─── The create writer ───────────────────────────────────────────────────────


def test_a_created_ad_carries_amazons_ad_id_and_a_refusal_carries_its_reason():
    body = {"productAds": {
        "success": [{"index": 0, "adId": "111"}, {"index": 2, "productAd": {"adId": "333"}},
                    {"index": 3}],
        "error": [{"index": 1, "errors": [{"errorType": "entityNotFoundError", "errorValue": {
            "entityNotFoundError": {"message": "no campaign found"}}}]}]}}
    out = {r["entity_id"]: r for r in spapi_ads._parse_create_outcome(body, ["p0", "p1", "p2", "p3", "p4"])}
    assert out["p0"]["ok"] and out["p0"]["ad_id"] == "111"
    assert out["p2"]["ad_id"] == "333"
    assert not out["p1"]["ok"] and "no campaign found" in out["p1"]["error"]
    assert not out["p3"]["ok"]                 # a success without an id cannot be undone
    assert not out["p4"]["ok"]                 # not mentioned is not done


async def test_the_create_request_is_amazons_seller_shape(monkeypatch):
    from app.config import get_settings
    monkeypatch.setattr(get_settings(), "ads_endpoint", "https://ads.test", raising=False)

    async def token(client):
        return "t"
    monkeypatch.setattr(spapi_ads, "_access_token", token)
    sent = {}

    class Response:
        status_code = 207

        def json(self):
            return {"productAds": {"success": [{"index": 0, "adId": "9"}], "error": []}}

    class Client:
        async def post(self, url, json, headers):
            sent.update(url=url, json=json, headers=headers)
            return Response()

    out = await spapi_ads.create_product_ads(Client(), [
        {"entity_id": "new-x", "campaign_id": "c1", "ad_group_id": "g1", "sku": "0.5kg cs 1"}])
    assert sent["url"] == "https://ads.test/sp/productAds"
    assert sent["json"] == {"productAds": [{"campaignId": "c1", "adGroupId": "g1", "sku": "0.5kg cs 1",
                                            "state": "ENABLED"}]}
    assert sent["headers"]["Content-Type"] == "application/vnd.spProductAd.v3+json"
    assert out == [{"entity_id": "new-x", "ok": True, "ad_id": "9", "error": None}]


# ─── The ledger ──────────────────────────────────────────────────────────────


async def test_a_created_ad_takes_its_real_id_and_its_undo_is_a_pause(db):
    changes = [{"entity_id": "new-a", "create": True, "writer": logic.WRITER_PRODUCT_AD, "text": "x",
                "ad_group_id": "g1", "campaign_id": "c1", "new_state": "ENABLED"},
               {"entity_id": "new-b", "create": True, "writer": logic.WRITER_PRODUCT_AD, "text": "y",
                "ad_group_id": "g1", "campaign_id": "c1", "new_state": "ENABLED"}]
    run = await repository.open_run(db, changes, rule_summary="SKU coverage: add 2 SKU(s)")
    rows = await repository.load_run(db, run)
    assert {r["status"] for r in rows} == {"pending"} and {r["action"] for r in rows} == {"create"}
    counts = await repository.record_create_results(db, run, [
        {"entity_id": "new-a", "ok": True, "ad_id": "555"},
        {"entity_id": "new-b", "ok": False, "ad_id": None, "error": "refused"}])
    assert counts == {"applied": 1, "failed": 1, "pending": 0}
    assert sorted(r["entity_id"] for r in await repository.load_run(db, run)) == ["555", "new-b"]
    undo = await repository.build_undo(db, run)
    assert [(u["entity_id"], u["old_state"], u["new_state"], u["writer"]) for u in undo] == [
        ("555", "ENABLED", "PAUSED", logic.WRITER_PRODUCT_AD)]


# ─── The routes ──────────────────────────────────────────────────────────────


@pytest.fixture
async def seeded(db, monkeypatch):
    from app.models import AdsEntity
    from app.portfolio import repository as pf_repo
    from app.shipment import catalogue as catalogue_mod

    for eid, c in CAMPAIGNS.items():
        db.add(AdsEntity(entity_type="campaign", ad_product="sp", entity_id=eid, name=c["name"],
                         state=c["state"]))
    for eid, g in GROUPS.items():
        db.add(AdsEntity(entity_type="ad_group", ad_product="sp", entity_id=eid, parent_id="c1",
                         name=g["name"], state=g["state"]))
    await db.commit()
    await repository.replace_product_ads(db, [
        {"entity_id": "1", "parent_id": "g1", "campaign_id": "c1", "name": "0.5kg cs 1 FBA",
         "asin": "A500", "state": "ENABLED"}])
    await listings.replace_listings(db, LISTINGS)

    async def load(**_):
        return {a: dict(r) for a, r in CATALOGUE.items()}, None, "sheet"

    async def flavours(_db, _cat):
        return {"PERI": "PERI", "NIMBU": "NIMBU"}

    async def parents(_db):
        # Amazon files both flavours under ONE parent; the flavour map splits them.
        return {"A500": "P-SATTU", "A1KG": "P-SATTU", "A2KG": "P-SATTU",
                "PERI": "P-CHANA", "NIMBU": "P-CHANA"}

    monkeypatch.setattr(catalogue_mod, "load_catalogue", load)
    monkeypatch.setattr(pf_repo, "flavour_keys", flavours)
    monkeypatch.setattr(pf_repo, "parent_of_children", parents)
    return db


@pytest.fixture
def amazon(monkeypatch):
    live = [{"entity_id": "1", "parent_id": "g1", "campaign_id": "c1", "name": "0.5kg cs 1 FBA",
             "sku": "0.5kg cs 1 FBA", "asin": "A500", "state": "ENABLED"}]
    created = []

    async def fetch_product_ads(client, *, ad_group_ids=None, states=()):
        assert ad_group_ids and states == ()
        return [a for a in live if a["parent_id"] in ad_group_ids]

    async def create(client, rows, **kw):
        created.extend(rows)
        return [{"entity_id": r["entity_id"], "ok": True, "ad_id": f"AD{i}", "error": None}
                for i, r in enumerate(rows)]

    monkeypatch.setattr(spapi_ads, "fetch_product_ads", fetch_product_ads)
    monkeypatch.setattr(spapi_ads, "create_product_ads", create)
    return live, created


async def test_the_panel_names_the_gaps(auth_client, seeded):
    body = (await auth_client.get("/ads/coverage")).json()
    assert body["totals"]["missing"] == len({
        "0.5kg cs 1", "0.5kg cs 1 flex", "1kg cs 1", "1kg cs 1 FBA"})
    assert body["listings"]["held"] == len(LISTINGS)


async def test_the_panel_splits_flavours_the_way_the_portfolio_does(auth_client, seeded):
    await repository.add_product_ads_to_cache(seeded, [
        {"entity_id": "9", "parent_id": "g2", "campaign_id": "c1", "name": "pp 250", "asin": "PERI"}])
    body = (await auth_client.get("/ads/coverage")).json()
    g2 = {m["sku"] for c in body["campaigns"] for g in c["groups"] if g["ad_group_id"] == "g2"
          for p in g["products"] for m in p["missing"]}
    assert g2 == {"pp 250 FBA"}


async def test_add_sends_only_current_gaps_and_never_a_sku_amazon_already_holds(auth_client, seeded, amazon):
    live, created = amazon
    live.append({"entity_id": "2", "parent_id": "g1", "campaign_id": "c1", "name": "1kg cs 1",
                 "sku": "1kg cs 1", "asin": "A1KG", "state": "PAUSED"})
    # Exactly what the screen posts: {rows: [{ad_group_id, sku}]}.
    body = (await auth_client.post("/ads/coverage/add", json={"rows": [
        {"ad_group_id": "g1", "sku": "0.5kg cs 1"},
        {"ad_group_id": "g1", "sku": "1kg cs 1"},        # added at Amazon since the morning
        {"ad_group_id": "g1", "sku": "0.5kg cs 1 FBA"},  # already in the group: not a gap
        {"ad_group_id": "g1", "sku": "2kg cs 1"},        # retired in the sheet: not a gap
    ]})).json()
    assert body["applied"] == 1, body
    assert [(r["ad_group_id"], r["sku"], r["campaign_id"]) for r in created] == [("g1", "0.5kg cs 1", "c1")]
    assert body["already_there"] == [{"ad_group_id": "g1", "sku": "1kg cs 1"}]
    assert {r["sku"] for r in body["not_a_gap"]} == {"0.5kg cs 1 FBA", "2kg cs 1"}

    rows = await repository.load_run(seeded, body["run_id"])
    assert [(r["entity_id"], r["status"], r["action"]) for r in rows] == [("AD0", "applied", "create")]
    after = (await auth_client.get("/ads/coverage")).json()
    still = {m["sku"] for c in after["campaigns"] for g in c["groups"] for p in g["products"]
             for m in p["missing"]}
    assert "0.5kg cs 1" not in still and "1kg cs 1" not in still


async def test_the_ledger_is_written_before_amazon_is_asked(auth_client, seeded, amazon, monkeypatch):
    seen = []

    async def create(client, rows, **kw):
        from app.database import async_session
        from app.models import AdsMutation
        from sqlalchemy import select
        async with async_session() as db:
            seen.extend((await db.execute(select(AdsMutation.status))).scalars().all())
        return [{"entity_id": r["entity_id"], "ok": True, "ad_id": "Z", "error": None} for r in rows]

    monkeypatch.setattr(spapi_ads, "create_product_ads", create)
    await auth_client.post("/ads/coverage/add", json={"rows": [{"ad_group_id": "g1", "sku": "0.5kg cs 1"}]})
    assert seen == ["pending"]


async def test_add_refuses_when_nothing_ticked_is_a_gap_or_over_the_row_limit(auth_client, seeded, amazon):
    _, created = amazon
    r = await auth_client.post("/ads/coverage/add", json={"rows": [{"ad_group_id": "g1", "sku": "nope"}]})
    assert r.status_code == 409
    await repository.save_guardrails(seeded, {**await repository.load_guardrails(seeded), "max_rows": 1})
    r = await auth_client.post("/ads/coverage/add", json={"rows": [
        {"ad_group_id": "g1", "sku": "0.5kg cs 1"}, {"ad_group_id": "g1", "sku": "1kg cs 1"}]})
    assert r.status_code == 400 and created == []


async def test_an_ignored_product_cannot_be_added_and_unignoring_restores_it(auth_client, seeded, amazon):
    await auth_client.post("/ads/coverage/ignore", json={"ad_group_id": "g1", "product_key": "P-SATTU"})
    body = (await auth_client.get("/ads/coverage")).json()
    assert body["totals"]["missing"] == 0 and body["totals"]["ignored"] == 4
    r = await auth_client.post("/ads/coverage/add", json={"rows": [{"ad_group_id": "g1", "sku": "0.5kg cs 1"}]})
    assert r.status_code == 409
    await auth_client.post("/ads/coverage/ignore",
                           json={"ad_group_id": "g1", "product_key": "P-SATTU", "ignored": False})
    assert (await auth_client.get("/ads/coverage")).json()["totals"]["missing"] == 4


async def test_undoing_an_add_pauses_the_new_ads(auth_client, seeded, amazon, monkeypatch):
    body = (await auth_client.post("/ads/coverage/add",
                                   json={"rows": [{"ad_group_id": "g1", "sku": "0.5kg cs 1"}]})).json()
    sent = []

    async def push(client, rows, *, writer):
        sent.extend((writer, r["entity_id"], r.get("new_state")) for r in rows)
        return [{"entity_id": r["entity_id"], "ok": True} for r in rows]
    monkeypatch.setattr(spapi_ads, "apply_changes", push)
    undo = (await auth_client.post(f"/ads/undo/{body['run_id']}")).json()
    assert undo["applied"] == 1
    assert sent == [(logic.WRITER_PRODUCT_AD, "AD0", "PAUSED")]


async def test_the_excel_lists_every_gap_on_screen(auth_client, seeded):
    import io

    from openpyxl import load_workbook
    r = await auth_client.get("/ads/coverage.xlsx")
    assert r.status_code == 200
    sheet = load_workbook(io.BytesIO(r.content)).active
    skus = {row[4] for row in sheet.iter_rows(values_only=True)}
    assert {"0.5kg cs 1", "0.5kg cs 1 flex", "1kg cs 1", "1kg cs 1 FBA"} <= skus


def test_the_run_history_names_an_add():
    from app.routers.ads import _change_words
    assert _change_words({"action": "create", "new_state": "ENABLED"}) == "added to the ad group (ON)"


# ─── The screen ──────────────────────────────────────────────────────────────


def test_the_screen_sends_only_the_ad_group_and_sku_pairs():
    body = PAGE[PAGE.index("async function addCoverage"):PAGE.index('addEventListener("input"')]
    assert "{ad_group_id: k.slice(0, at), sku: k.slice(at + 1)}" in body
    assert 'fetch("/ads/coverage/add"' in body and "JSON.stringify({rows})" in body


def test_the_screen_never_ticks_an_ignored_row():
    rows = PAGE[PAGE.index("function coverageRowsHtml"):PAGE.index("function renderCoverage")]
    assert 'product.ignored ? "" : `<input type="checkbox" data-cov-pick' in rows
    group = PAGE[PAGE.index('const group = e.target.closest("[data-cov-group]")'):]
    assert "if(!p.ignored) for(const m of p.missing)" in group[:600]
    # The card's own "select all" tick reads the same rule.
    keys = PAGE[PAGE.index("function groupKeys"):PAGE.index("function coverageRowsHtml")]
    assert "g.products.filter(p => !p.ignored)" in keys


# ─── The SKU coverage sub-tab (first named Expansion) ───────────────────────────────────────────────────


async def test_expansion_is_its_own_ads_sub_tab_with_one_card_per_ad_group(auth_client):
    r = await auth_client.get("/ads-page/expansion")
    assert r.status_code == 200
    html = r.text
    assert 'href="/ads-page">Optimisation</a>' in html
    assert 'class="seg-btn on" href="/ads-page/expansion">SKU coverage</a>' in html
    assert 'id="coverage-cards"' in html and "cov-card" in html


async def test_optimisation_carries_the_tabs_and_no_longer_the_coverage_panel(auth_client):
    html = (await auth_client.get("/ads-page")).text
    assert 'class="seg-btn on" href="/ads-page">Optimisation</a>' in html
    assert 'href="/ads-page/expansion">SKU coverage</a>' in html
    assert 'href="/ads-page/growth">Keyword growth</a>' in html
    assert "coverage-panel" not in OPTIMISATION and "coverage-btn" not in OPTIMISATION
    assert "/ads/coverage" not in OPTIMISATION


def test_a_card_is_one_ad_group():
    cards = PAGE[PAGE.index("function coverageGroups"):PAGE.index("function groupKeys")]
    # One entry per (campaign, ad group), never one per product or per SKU.
    assert "for(const c of (coverageData && coverageData.campaigns) || []) for(const g of c.groups){" in cards
    assert "out.push({c, g, products});" in cards
    render = PAGE[PAGE.index("function renderCoverage"):PAGE.index("async function undoRun")]
    assert "groups.map(cardHtml)" in render


async def test_the_expansion_page_needs_the_ads_area(db_schema):
    from httpx import ASGITransport, AsyncClient

    from app.main import app
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get("/ads-page/expansion", follow_redirects=False)
    assert r.status_code in (302, 303, 307) and "/login" in r.headers["location"]


def test_a_big_card_shows_eight_rows_and_a_show_all_link():
    card = PAGE[PAGE.index("function cardHtml"):PAGE.index("function renderCoverage")]
    assert "const CARD_ROWS = 8;" in PAGE
    assert "let left = open ? Infinity : CARD_ROWS;" in card
    assert "coverageRowsHtml(g, p, left)" in card
    assert "product.missing.slice(0, limit)" in PAGE
    assert "total > CARD_ROWS" in card and "data-cov-open" in card


def test_every_row_of_a_card_is_the_same_grid_so_the_columns_line_up():
    assert ".cov-grid-row{display:grid;grid-template-columns:" in PAGE
    rows = PAGE[PAGE.index("function coverageRowsHtml"):PAGE.index("function cardHtml")]
    assert 'class="cov-grid-row cov-row"' in rows
    card = PAGE[PAGE.index("function cardHtml"):PAGE.index("function renderCoverage")]
    assert 'class="cov-grid-row cov-cols"' in card and 'class="cov-grid-row cov-ptitle' in card
