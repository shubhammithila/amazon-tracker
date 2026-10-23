"""The 7d / 14d / 30d buttons must be quick (23 Oct 2026).

Reported as *"when I click from 7d to 30d or 14d the transition is very slow"*. Measured on
production: GET /ads took 2.0 / 2.6 / 3.8 s, of which `sum_daily` was 0.5 / 1.1 / 2.4 s — every keyword
row of the window fetched only to make 29 campaign totals — and ~1 s more was ten separate scans
asking which days are held. The totals are now summed in SQL and the held days read once.
"""
from app.ads import repository


def _rows(day, campaign, n, product="sp"):
    return [{"date": day, "keywordId": f"{campaign}-{product}-{i}", "matchType": "EXACT",
             "keyword": f"kw{i}", "campaignId": campaign, "adGroupId": "g1",
             "cost": 10.0 + i, "sales7d": 5.0 * i, "purchases7d": i % 2, "clicks": i,
             "impressions": 10 * i, "keywordBid": 5.0, "sales": 5.0 * i, "purchases": i % 2}
            for i in range(n)]


async def _seed(db):
    for day in ("2026-10-20", "2026-10-21", "2026-10-22"):
        await repository.save_daily(db, _rows(day, "c1", 4) + _rows(day, "c2", 3), ad_product="sp")
        await repository.save_daily(db, _rows(day, "c3", 2, "sb"), ad_product="sb")
    await repository.save_entities(db, [
        {"entity_type": "campaign", "entity_id": c, "ad_product": p, "name": f"camp {c}", "state": "ENABLED"}
        for c, p in (("c1", "sp"), ("c2", "sp"), ("c3", "sb"))])


async def test_campaign_totals_equal_the_old_rollup_of_every_keyword_row(db):
    await _seed(db)
    rows = await repository.sum_daily(db, "2026-10-21", "2026-10-22")
    expected: dict = {}
    for r in rows:
        b = expected.setdefault(r["campaign_id"], {"spend": 0.0, "sales": 0.0, "clicks": 0,
                                                    "impressions": 0, "orders": 0, "rows": 0})
        for k in ("spend", "sales", "clicks", "impressions", "orders"):
            b[k] += r[k]
        b["rows"] += 1
    got = await repository.campaign_totals(db, "2026-10-21", "2026-10-22")
    assert set(got) == set(expected) == {"c1", "c2", "c3"}
    for cid, b in expected.items():
        for k, v in b.items():
            assert round(got[cid][k], 2) == round(v, 2), (cid, k)
    only_sp = await repository.campaign_totals(db, "2026-10-21", "2026-10-22", ad_product="sp")
    assert set(only_sp) == {"c1", "c2"}


async def test_the_held_days_are_read_once_and_give_the_same_answer(db):
    await _seed(db)
    held = await repository.held_days_by_product(db)
    assert held == {"sp": {"2026-10-20", "2026-10-21", "2026-10-22"},
                    "sb": {"2026-10-20", "2026-10-21", "2026-10-22"}}
    for start, end in (("2026-10-20", "2026-10-22"), ("2026-10-19", "2026-10-22")):
        assert (await repository.range_completeness(db, start, end, held=held)
                == await repository.range_completeness(db, start, end))


async def test_the_campaign_list_does_not_read_every_keyword_row(auth_client, db, monkeypatch):
    await _seed(db)
    calls = []
    original = repository.sum_daily

    async def counting(*a, **kw):
        calls.append(kw.get("campaign_ids"))
        return await original(*a, **kw)

    monkeypatch.setattr(repository, "sum_daily", counting)
    body = (await auth_client.get("/ads?start=2026-10-20&end=2026-10-22")).json()
    assert calls == [], "the campaign list summed every keyword row again"
    assert body["totals"]["spend"] == round(3 * (sum(10.0 + i for i in range(4))
                                                 + sum(10.0 + i for i in range(3))
                                                 + sum(10.0 + i for i in range(2))), 2)
    await auth_client.get("/ads?start=2026-10-20&end=2026-10-22&campaign_id=c1")
    assert calls == [["c1"]]


async def test_the_window_note_still_counts_every_keyword_row(auth_client, db):
    """"27728 target rows cached" under the window read 0 once the rows stopped being fetched."""
    await _seed(db)
    body = (await auth_client.get("/ads?start=2026-10-20&end=2026-10-22")).json()
    assert body["totals"]["targets"] == 4 + 3 + 2


# ─── Run history: the Excel download beside Undo ────────────────────────────


async def test_a_run_downloads_as_excel_with_every_change_and_its_result(auth_client, db):
    import io
    from datetime import datetime

    from openpyxl import load_workbook

    from app.models import AdsMutation
    await repository.save_entities(db, [
        {"entity_type": "campaign", "entity_id": "c1", "name": "MF_SP_keywords", "state": "ENABLED"},
        {"entity_type": "ad_group", "entity_id": "g1", "name": "chana sattu", "state": "ENABLED"}])
    when = datetime(2026, 10, 23, 12, 18)            # 17:48 IST
    for i, (old, new, state, status, err) in enumerate([
            (10.0, 9.0, None, "applied", None), (None, None, "PAUSED", "applied", None),
            (5.0, 4.5, None, "failed", "rangeError: too low")]):
        db.add(AdsMutation(run_id="run-1", entity_id=f"E{i}", entity_type="keyword", writer="keyword",
                           text=f"kw{i}", campaign_id="c1", ad_group_id="g1", old_bid=old, new_bid=new,
                           old_state="ENABLED" if state else None, new_state=state,
                           action="state" if state else "bid", status=status, error=err,
                           rule_summary="Spend > 100 -> bid decrease 10%", created_at=when))
    await db.commit()
    r = await auth_client.get("/ads/runs/run-1/download.xlsx")
    assert r.status_code == 200
    assert "ads-run-2026-10-23-1748.xlsx" in r.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(r.content)).active
    cells = [[c.value for c in row] for row in sheet.iter_rows()]
    text = " ".join(str(v) for row in cells for v in row if v is not None)
    assert "17:48 IST" in text and "Spend > 100 -> bid decrease 10%" in text
    body = [row for row in cells if row and row[2] in ("kw0", "kw1", "kw2")]
    assert [row[0] for row in body] == ["MF_SP_keywords"] * 3 and body[0][1] == "chana sattu"
    assert body[0][5] == "bid ₹10.00 → ₹9.00" and body[1][5] == "turned OFF"
    assert body[2][10] == "refused by Amazon" and body[2][11] == "rangeError: too low"
    assert (await auth_client.get("/ads/runs/nope/download.xlsx")).status_code == 404


def test_run_history_offers_the_download_and_reads_ist():
    from pathlib import Path
    page = (Path(__file__).resolve().parent.parent / "templates" / "ads.html").read_text(encoding="utf-8")
    assert "/download.xlsx" in page and "⬇ Excel" in page
    assert "esc(istStamp(r.at))" in page
