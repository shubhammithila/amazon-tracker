"""Portfolio -> Analytics: two years of monthly history for the active range (17 Oct 2026)."""
import ast
from pathlib import Path

import pytest
from sqlalchemy import func, select

from app.analytics import fetch, logic, repository
from app.database import async_session
from app.models import AnalyticsMonthly, EconomicsDaily

pytestmark = pytest.mark.regression
ROOT = Path(__file__).resolve().parent.parent


# ── pure arithmetic ─────────────────────────────────────────────────────────────────────────────

def test_month_range_crosses_a_year():
    assert logic.month_range("2024-11", "2025-02") == ["2024-11", "2024-12", "2025-01", "2025-02"]


def size(sales, units, kg):
    return {"sales": sales, "units": units, "kg": kg}


def test_the_growth_split_adds_up_exactly_with_a_pack_in_only_one_month():
    prev = {"A": size(1000.0, 10, 5.0), "B": size(900.0, 3, 6.0)}
    cur = {"A": size(1320.0, 12, 6.0), "C": size(450.0, 1, 2.0)}
    s = logic.growth_split(prev, cur)
    assert s["delta"] == pytest.approx(1770 - 1900)
    assert s["volume"] + s["price"] + s["mix"] == pytest.approx(s["delta"], abs=0.01)
    # A: 100 -> 110 a pack, 12 packs: +120 of price, the only pack in both months.
    assert s["price"] == pytest.approx(120.0)
    # kg 11 -> 8 at last month's 1900/11 per kg.
    assert s["volume"] == pytest.approx(-3 * 1900 / 11, abs=0.01)


def test_a_price_rise_on_an_unchanged_pack_is_price_only():
    s = logic.growth_split({"A": size(1000.0, 10, 5.0)}, {"A": size(1100.0, 10, 5.0)})
    assert s == {"delta": 100.0, "volume": 0.0, "price": 100.0, "mix": 0.0}


def test_a_launch_is_all_volume():
    s = logic.growth_split({}, {"A": size(500.0, 5, 2.5)})
    assert s == {"delta": 500.0, "volume": 500.0, "price": 0.0, "mix": 0.0}


def test_a_shift_to_cheaper_kilograms_shows_as_mix():
    """Same kg and same pack prices, but more of the cheaper-per-kg big pack: mix, not price."""
    prev = {"small": size(600.0, 6, 3.0), "big": size(200.0, 2, 2.0)}       # 200/kg and 100/kg
    cur = {"small": size(400.0, 4, 2.0), "big": size(300.0, 3, 3.0)}         # same 5 kg, cheaper mix
    s = logic.growth_split(prev, cur)
    assert s["price"] == 0.0 and s["volume"] == 0.0
    assert s["mix"] == pytest.approx(s["delta"]) and s["mix"] < 0


def test_an_unknown_weight_falls_back_to_units():
    s = logic.growth_split({"A": size(1000.0, 10, None)}, {"A": size(1500.0, 15, None)})
    assert s["volume"] == pytest.approx(500.0) and s["price"] == 0.0


def test_concentration_is_the_top_n_share():
    sales = {f"p{i}": float(i) for i in range(1, 21)}           # 1..20, total 210
    assert logic.concentration(sales, 10) == pytest.approx(sum(range(11, 21)) / 210)
    assert logic.concentration({}, 10) is None


def test_conversion_comes_from_sums_not_an_average_of_rates():
    rows = [{"sessions": 100, "page_views": 120, "units": 10, "buy_box_pct": 100.0},
            {"sessions": 900, "page_views": 1000, "units": 9, "buy_box_pct": 50.0}]
    r = logic.weighted_rates(rows)
    assert r["conversion"] == pytest.approx(19 / 1000)               # a plain average says 5.5%
    assert r["buy_box"] == pytest.approx((100 * 120 + 50 * 1000) / 1120 / 100)
    assert logic.weighted_rates([]) == {"sessions": None, "page_views": None,
                                        "conversion": None, "buy_box": None}


def test_growth_is_none_without_a_previous_month():
    assert logic.growth(100.0, 0) is None and logic.growth(110.0, 100.0) == pytest.approx(0.1)


# ── fetch and store ─────────────────────────────────────────────────────────────────────────────

def raw(month, asin, sales, units, *, parent="B0PARENT01", fee=100.0, fee_tax=18.0, ads=50.0, net=600.0):
    return {"startDate": f"{month}-01", "parentAsin": parent, "childAsin": asin,
            "sales": {"orderedProductSales": {"amount": sales}, "refundedProductSales": {"amount": 0},
                      "unitsOrdered": units, "unitsRefunded": 0, "netUnitsSold": units},
            "fees": [{"feeTypeName": "Referral", "charges": [{"aggregatedDetail": {
                "totalAmount": {"amount": fee}, "taxAmount": {"amount": fee_tax}}}]}],
            "ads": [{"adTypeName": "SponsoredProductFee", "charge": {"totalAmount": {"amount": ads}}}],
            "netProceeds": {"total": {"amount": net}}}


def test_a_stored_month_is_written_through_the_profit_tabs_rule():
    """Fees ex-GST and the fee GST added back to Net, exactly as `size_row` stores a Profit day."""
    row = repository.rows_from_economics([raw("2025-10", "B0CHILD001", 1000.0, 10)])[0]
    assert row["month"] == "2025-10" and row["ordered_sales"] == 1000.0
    assert row["net_proceeds"] == pytest.approx(618.0)                 # 600 + 18 GST
    assert '"Referral": 82.0' in row["fees_json"]
    assert repository.rows_from_economics([{**raw("2025-10", "X", 1.0, 1), "startDate": None}]) == []


async def test_the_economics_query_asks_for_months():
    seen = {}

    async def run(query, label):
        seen["q"] = query
        return []
    await fetch.economics_months("2025-01-01", "2025-03-31", run=run)
    assert "date: MONTH" in seen["q"] and "date: RANGE" not in seen["q"]


def test_the_current_month_ends_yesterday():
    from datetime import date
    assert fetch.month_bounds("2026-10", date(2026, 10, 17)) == ("2026-10-01", "2026-10-16")
    assert fetch.month_bounds("2026-10", date(2026, 10, 1)) is None
    assert fetch.month_bounds("2026-09", date(2026, 10, 17)) == ("2026-09-01", "2026-09-30")


async def test_saving_a_month_replaces_it_and_keeps_the_others(db_schema):
    async with async_session() as db:
        await repository.save_economics(db, repository.rows_from_economics(
            [raw("2025-09", "B0A", 100.0, 1), raw("2025-10", "B0A", 200.0, 2)]), ["2025-09", "2025-10"])
        await repository.save_economics(db, repository.rows_from_economics(
            [raw("2025-10", "B0A", 300.0, 3)]), ["2025-10"])
        await db.commit()
        rows = {(r.month, float(r.ordered_sales)) for r in (await db.execute(select(AnalyticsMonthly))).scalars()}
    assert rows == {("2025-09", 100.0), ("2025-10", 300.0)}


async def test_sponsored_brands_is_added_to_ads_and_taken_off_net_where_held(db_schema):
    async with async_session() as db:
        await repository.save_economics(db, repository.rows_from_economics(
            [raw("2026-09", "B0A", 1000.0, 10)]), ["2026-09"])
        db.add(EconomicsDaily(day="2026-09-05", child_asin="B0A", seller_sku="", sb_spend=40,
                              sb_basis="declared", ordered_sales=0, net_proceeds=0))
        await db.commit()
        months = await repository.load_months(db)
    row = months["2026-09"][0]
    ads = {a["adTypeName"]: a["charge"]["totalAmount"]["amount"] for a in row["ads"]}
    assert sum(ads.values()) == pytest.approx(90.0)                    # 50 SP + 40 SB
    assert row["netProceeds"]["total"]["amount"] == pytest.approx(618.0 - 40)


@pytest.mark.parametrize("module", ["fetch", "repository", "logic"])
def test_analytics_never_writes_to_amazon(module):
    tree = ast.parse((ROOT / "app" / "analytics" / f"{module}.py").read_text(encoding="utf-8"))
    names = ({n.attr for n in ast.walk(tree) if isinstance(n, ast.Attribute)}
             | {a.name for n in ast.walk(tree) if isinstance(n, ast.ImportFrom) for a in n.names})
    assert not names & {"_post", "save_economics_daily", "apply_changes", "create_inbound_plan"}


# ── the payload ─────────────────────────────────────────────────────────────────────────────────

from datetime import date
from types import SimpleNamespace

from app.analytics import service

CATALOGUE = {
    "B0CS500": {"name": "Chana Sattu", "weight": 0.5, "brand": "Mithila Foods", "active": True},
    "B0CS1KG": {"name": "Chana Sattu", "weight": 1.0, "brand": "Mithila Foods", "active": True},
    "B0RICE1": {"name": "Usna Chawal", "weight": 1.0, "brand": "Mithila Foods", "active": True},
    "B0OLD01": {"name": "Old Thekua", "weight": 0.35, "brand": "Mithila Foods", "active": False},
}


@pytest.fixture
async def stored(db_schema, monkeypatch):
    service._CACHE.clear()

    async def cat():
        return CATALOGUE, None, "sheet"

    async def cats(db):
        return [SimpleNamespace(product_key="chana sattu", priority=1),
                SimpleNamespace(product_key="usna chawal", priority=4)]
    monkeypatch.setattr("app.shipment.catalogue.load_catalogue", cat)
    monkeypatch.setattr("app.shipment.repository.load_categories", cats)
    data = [
        # Aug: chana 500g 10 x 100, rice 5 x 200; old thekua sells
        raw("2026-08", "B0CS500", 1000.0, 10, parent="B0CSPAR", net=500.0, fee_tax=0),
        raw("2026-08", "B0RICE1", 1000.0, 5, parent="B0RICE1", net=400.0, fee_tax=0),
        raw("2026-08", "B0OLD01", 300.0, 3, parent="B0OLD01", net=50.0, fee_tax=0),
        # Sep: chana 500g price up to 110, plus 1 kg packs; rice flat
        raw("2026-09", "B0CS500", 1320.0, 12, parent="B0CSPAR", net=700.0, fee_tax=0),
        raw("2026-09", "B0CS1KG", 400.0, 2, parent="B0CSPAR", net=200.0, fee_tax=0),
        raw("2026-09", "B0RICE1", 1000.0, 5, parent="B0RICE1", net=400.0, fee_tax=0),
        # Oct (month to date)
        raw("2026-10", "B0CS500", 330.0, 3, parent="B0CSPAR", net=150.0, fee_tax=0),
    ]
    async with async_session() as db:
        rows = repository.rows_from_economics(data)
        await repository.save_economics(db, rows, ["2026-08", "2026-09", "2026-10"])
        await repository.save_traffic(db, repository.rows_from_traffic([
            {"childAsin": "B0CS500", "parentAsin": "B0CSPAR", "sales": {"unitsOrdered": 12},
             "traffic": {"sessions": 100, "pageViews": 120, "buyBoxPercentage": 90.0}},
            {"childAsin": "B0CS1KG", "parentAsin": "B0CSPAR", "sales": {"unitsOrdered": 2},
             "traffic": {"sessions": 100, "pageViews": 80, "buyBoxPercentage": 100.0}}], "2026-09"), "2026-09")
        await db.commit()
    yield
    service._CACHE.clear()


TODAY = date(2026, 10, 17)


async def test_inactive_products_are_hidden_until_asked(stored):
    async with async_session() as db:
        shown = await service.build(db, today=TODAY)
        everything = await service.build(db, include_inactive=True, today=TODAY)
    assert {p["product"] for p in shown["products"]} == {"Chana Sattu", "Usna Chawal"}
    assert "Old Thekua" in {p["product"] for p in everything["products"]}
    assert shown["totals"]["2026-08"]["sales"] == 2000.0
    assert everything["totals"]["2026-08"]["sales"] == 2300.0


async def test_the_current_month_is_month_to_date_and_never_a_growth_figure(stored):
    async with async_session() as db:
        data = await service.build(db, today=TODAY)
    assert data["complete_months"] == ["2026-08", "2026-09"] and data["last_month"] == "2026-09"
    assert data["mtd"] == {"month": "2026-10", "days": 16}
    assert data["totals"]["2026-10"]["mom"] is None and "2026-10" not in data["growth"]
    assert data["totals"]["2026-09"]["mom"] == pytest.approx(2720 / 2000 - 1)


async def test_totals_are_the_sum_of_the_products_and_the_split_adds_up(stored):
    async with async_session() as db:
        data = await service.build(db, today=TODAY)
    sep = data["totals"]["2026-09"]
    assert sep["sales"] == pytest.approx(sum(p["series"]["2026-09"]["sales"] for p in data["products"]))
    assert sep["net_pct"] == pytest.approx(1300 / 2720)
    g = data["growth"]["2026-09"]
    assert g["delta"] == pytest.approx(720.0)
    assert g["volume"] + g["price"] + g["mix"] == pytest.approx(g["delta"], abs=0.02)
    assert g["price"] == pytest.approx(120.0)                      # 100 -> 110 on 12 packs


async def test_categories_are_the_shipment_tabs_and_sum_their_rows(stored):
    async with async_session() as db:
        data = await service.build(db, today=TODAY)
    by = {c["category"]: c for c in data["categories"]}
    assert by["Sattu"]["months"]["2026-09"]["sales"] == pytest.approx(1720.0)
    assert by["Rice"]["months"]["2026-09"]["share"] == pytest.approx(1000 / 2720)
    async with async_session() as db:
        only = await service.build(db, category="Rice", today=TODAY)
    assert [p["product"] for p in only["products"]] == ["Usna Chawal"]


async def test_a_product_row_carries_price_traffic_and_net(stored):
    async with async_session() as db:
        data = await service.build(db, today=TODAY)
    chana = next(p for p in data["products"] if p["product"] == "Chana Sattu")
    assert chana["price_pack"] == "500g" and chana["price_now"] == pytest.approx(110.0)
    assert chana["price_low"] == pytest.approx(100.0) and chana["price_high"] == pytest.approx(110.0)
    assert chana["sessions"] == 200 and chana["conversion"] == pytest.approx(14 / 200)
    assert chana["buy_box"] == pytest.approx((90 * 120 + 100 * 80) / 200 / 100)
    assert chana["mom"] == pytest.approx(1720 / 1000 - 1)
    rice = next(p for p in data["products"] if p["product"] == "Usna Chawal")
    assert rice["sessions"] is None                                   # no traffic row: a dash


async def test_a_month_is_the_profit_tabs_portfolio_of_the_same_rows(stored):
    from app.portfolio import logic as pf_logic
    async with async_session() as db:
        months = await repository.load_months(db)
        data = await service.build(db, today=TODAY)
    direct = pf_logic.portfolio(months["2026-09"], CATALOGUE, {}, {}, flavour_keys={})
    assert data["totals"]["2026-09"]["net"] == pytest.approx(direct["totals"]["net"])
    assert data["totals"]["2026-09"]["sales"] == pytest.approx(direct["totals"]["sales"])


# ── routes, refresh and the page ────────────────────────────────────────────────────────────────

async def test_the_route_returns_the_payload_and_needs_the_portfolio_area(stored, auth_client, ops_client):
    r = await auth_client.get("/portfolio/analytics?category=Rice")
    assert r.status_code == 200 and r.json()["category"] == "Rice"
    assert (await ops_client.get("/portfolio/analytics")).status_code in (303, 403)


async def test_a_refresh_while_one_runs_is_409(auth_client):
    from app.analytics import refresh
    refresh.STATE["running"] = True
    try:
        assert (await auth_client.post("/portfolio/analytics/refresh")).status_code == 409
    finally:
        refresh.STATE["running"] = False


@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
async def test_the_download_carries_the_products_table(stored, auth_client, fmt):
    r = await auth_client.post("/portfolio/analytics/export", json={"format": fmt})
    assert r.status_code == 200 and len(r.content) > 500
    if fmt == "xlsx":
        import io
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        assert "Price now" in [c.value for c in ws[2]]


async def test_a_refresh_stores_months_and_a_failure_stops_and_is_recorded(db_schema):
    from app.analytics import refresh

    async def run(query, label):
        if "traffic" in label:
            return [{"childAsin": "B0A", "parentAsin": "B0P", "sales": {"unitsOrdered": 1},
                     "traffic": {"sessions": 10, "pageViews": 12, "buyBoxPercentage": 99.0}}]
        return [raw("2026-08", "B0A", 100.0, 1), raw("2026-09", "B0A", 200.0, 2)]
    out = await refresh.run(["2026-08", "2026-09"], today=TODAY, run_query=run)
    assert out["status"] == "done"
    async with async_session() as db:
        held = await repository.months_held(db)
        assert held == {"economics": {"2026-08", "2026-09"}, "traffic": {"2026-08", "2026-09"}}

    async def boom(query, label):
        raise RuntimeError("Amazon FATAL")
    out = await refresh.run(["2026-10"], today=TODAY, run_query=boom)
    assert out["status"] == "failed" and "FATAL" in out["error"] and not refresh.STATE["running"]
    async with async_session() as db:
        assert (await repository.last_run(db))["status"] == "failed"


def test_the_nightly_months_are_this_month_and_last():
    from app.analytics import refresh
    assert refresh.recent_months(date(2026, 10, 17)) == ["2026-09", "2026-10"]
    assert refresh.recent_months(date(2026, 11, 1)) == ["2026-10"]
    assert refresh.recent_months(date(2027, 1, 5)) == ["2026-12", "2027-01"]


PAGE = (ROOT / "templates" / "portfolio_analytics.html").read_text(encoding="utf-8")


def test_it_is_a_portfolio_sub_tab():
    tabs = (ROOT / "templates" / "_portfolio_tabs.html").read_text(encoding="utf-8")
    assert 'href="/portfolio-page/analytics"' in tabs and '{% include "_portfolio_tabs.html" %}' in PAGE


async def test_the_page_opens(auth_client):
    r = await auth_client.get("/portfolio-page/analytics")
    assert r.status_code == 200 and 'id="mom-chart"' in r.text


def test_the_month_to_date_is_drawn_apart_and_charts_read_colour_from_the_theme():
    assert '" (to date)"' in PAGE and "themeColour(" in PAGE
    assert "getComputedStyle(document.documentElement)" in PAGE


def test_a_month_holding_a_pack_with_no_known_weight_has_no_kg_total():
    """A dash, never a short sum: 10 kg known plus an unknown pack is not "10 kg"."""
    known = {"sales": 100.0, "units": 10, "kg": 10.0, "net": 1.0, "fees": 1.0, "refunded": 0, "ad_spend": 0}
    unknown = {**known, "kg": None}
    assert service._month_totals([known])["kg"] == 10.0
    assert service._month_totals([known, unknown])["kg"] is None


async def test_a_quota_refusal_waits_and_retries(monkeypatch):
    waits, calls = [], {"n": 0}

    async def sleep(s):
        waits.append(s)

    async def query(q, label, sleep):
        calls["n"] += 1
        if calls["n"] < 3:
            raise RuntimeError("Amazon said: You exceeded your quota for the requested resource.")
        return [{"ok": 1}]
    monkeypatch.setattr("app.portfolio.economics._run_query", query)
    assert await fetch._default_run("q", "label", sleep=sleep) == [{"ok": 1}]
    assert waits == list(fetch.QUOTA_WAITS[:2])


async def test_any_other_failure_is_not_retried(monkeypatch):
    async def query(q, label, sleep):
        raise RuntimeError("Amazon FATAL")

    async def sleep(s):
        raise AssertionError("must not wait")
    monkeypatch.setattr("app.portfolio.economics._run_query", query)
    with pytest.raises(RuntimeError, match="FATAL"):
        await fetch._default_run("q", "label", sleep=sleep)


async def test_a_month_missing_its_traffic_is_still_to_do(db_schema):
    from app.analytics import refresh
    async with async_session() as db:
        await repository.save_economics(db, repository.rows_from_economics([raw("2024-11", "B0A", 1.0, 1)]), ["2024-11"])
        await db.commit()
    missing = await refresh.missing_months(date(2025, 1, 10))
    assert missing == ["2024-11", "2024-12"]


async def test_a_part_month_is_asked_for_as_a_range_because_MONTH_drops_it():
    """Measured 18 Oct 2026: MONTH over 1 Sep..17 Oct returned September alone."""
    asked = []

    async def run(query, label):
        asked.append(("MONTH" if "date: MONTH" in query else "RANGE", label))
        return []
    await fetch.economics_months("2026-09-01", "2026-10-17", run=run)
    assert [a[0] for a in asked] == ["MONTH", "RANGE"]
    assert "2026-10-01..2026-10-17" in asked[1][1] and "2026-09-01..2026-09-30" in asked[0][1]
    asked.clear()
    await fetch.economics_months("2026-10-01", "2026-10-17", run=run)
    assert [a[0] for a in asked] == ["RANGE"]
    asked.clear()
    await fetch.economics_months("2026-07-01", "2026-09-30", run=run)
    assert [a[0] for a in asked] == ["MONTH"]


def test_a_child_split_across_two_parents_is_summed_not_overwritten():
    """Aug 2026: 249 traffic rows for 129 ASINs, a listing moved to a new parent mid-month."""
    rows = repository.rows_from_traffic([
        {"childAsin": "B0CWGXYLT6", "parentAsin": "B0CWGXYLT6", "sales": {"unitsOrdered": 213, "orderedProductSales": {"amount": 39406.44}},
         "traffic": {"sessions": 1185, "pageViews": 1467, "buyBoxPercentage": 99.93}},
        {"childAsin": "B0CWGXYLT6", "parentAsin": "B0DR322QYG", "sales": {"unitsOrdered": 161, "orderedProductSales": {"amount": 29240.04}},
         "traffic": {"sessions": 731, "pageViews": 1032, "buyBoxPercentage": 99.89}}], "2026-08")
    assert len(rows) == 1
    r = rows[0]
    assert r["sessions"] == 1916 and r["page_views"] == 2499 and r["units_ordered"] == 374
    assert r["sales"] == pytest.approx(68646.48) and r["parent_asin"] == "B0CWGXYLT6"
    assert r["buy_box_pct"] == pytest.approx((99.93 * 1467 + 99.89 * 1032) / 2499, abs=0.01)


def test_the_growth_split_table_is_gone_and_categories_are_a_chart():
    """Asked for: remove "Why sales moved", and categories by month as a graph, not a grid."""
    assert "Why sales moved" not in PAGE and 'id="growth"' not in PAGE and "renderGrowth" not in PAGE
    assert 'id="cat-grid"' not in PAGE and 'id="cat-chart"' in PAGE and "function renderCatChart(" in PAGE


def test_every_trend_chart_has_a_view_switch_a_readout_and_follows_the_range():
    for seg, read in (("cat-view", "cat-read"), ("cost-view", "cost-read"),
                      ("funnel-view", "funnel-read"), ("kg-view", "kg-read")):
        assert f'id="{seg}"' in PAGE and f'id="{read}"' in PAGE
    for fn in ("renderCatChart", "renderCostChart", "renderFunnelChart", "renderKgChart"):
        body = PAGE[PAGE.index(f"function {fn}("):]
        body = body[:body.index("\nfunction ")]
        assert "chartMonths(" in body, fn
    # hovering a month shows every series for that month
    assert 'interaction = {mode: "index", intersect: false}' in PAGE


async def test_per_kg_tacos_and_category_traffic_come_from_sums(stored):
    async with async_session() as db:
        data = await service.build(db, today=TODAY)
    sep = data["totals"]["2026-09"]
    assert sep["per_kg"] == pytest.approx(sep["sales"] / sep["kg"])
    sattu = next(c for c in data["categories"] if c["category"] == "Sattu")["months"]["2026-09"]
    assert sattu["per_kg"] == pytest.approx(sattu["sales"] / sattu["kg"])
    assert sattu["sessions"] == 200 and sattu["conversion"] == pytest.approx(14 / 200)
    assert sattu["buy_box"] == pytest.approx((90 * 120 + 100 * 80) / 200 / 100)
    chana = next(p for p in data["products"] if p["product"] == "Chana Sattu")["series"]["2026-09"]
    assert chana["per_kg"] == pytest.approx(chana["sales"] / chana["kg"])
    assert chana["tacos"] == (chana["ad_spend"] / chana["sales"] if chana["ad_spend"] is not None else None)


def test_per_kg_is_a_dash_when_a_pack_weight_is_unknown():
    known = {"sales": 100.0, "units": 10, "kg": 10.0, "net": 1.0, "fees": 1.0, "refunded": 0, "ad_spend": 0}
    assert service._month_totals([known])["per_kg"] == pytest.approx(10.0)
    assert service._month_totals([known, {**known, "kg": None}])["per_kg"] is None


async def test_every_category_is_named_even_when_one_is_picked(stored):
    async with async_session() as db:
        only = await service.build(db, category="Rice", today=TODAY)
    assert set(only["category_names"]) == {"Sattu", "Rice"}
