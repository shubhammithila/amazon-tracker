"""The MRP sheet's Active flag decides which products the Portfolio tab counts.

Asked for as *"whether the sku/parent is killed or not can be taken from the MRP sheet… in which
column V has the data in terms of Y and N. Refresh it every day."*

The column was already parsed — `app/shipment/catalogue.py` has read it for the Shipment tab all
along — and `_dashboard` already loaded the catalogue. What was missing was a CONSUMER: `size_row`
read name, brand and weight off each entry and ignored `active`.

**Measured on the live 30-day window before building any of this**, because the scale of the change
is the thing that needed deciding rather than the mechanism:

* 157 of 267 SKUs are marked N, and **55 of 90 parents are inactive in every size** — so the tab
  goes from 90 products to 35 and the Sales KPI drops Rs 67,193;
* **11 of those SKUs, across 8 parents, still SOLD** — Moringa Powder 70u, Bengali Moori 121u,
  Bengali Banskathi 20u, Makhana Powder 15u, Bengali Miniket 11u, Herbal Gulal 3u, Moori 4u,
  Peanut Thekua 1u;
* 6 parents are MIXED, and all six have zero sales on their inactive sizes **today** — so live data
  cannot exercise the case that matters and the fixtures here construct it;
* 0 of 267 selling SKUs are absent from the sheet, so the unknown-ASIN default is likewise
  untestable against production and is pinned below.

That 1.5% gap is the **3,337-vs-3,259** report in a different tab: an inactive product that is still
selling is a question, not a fact — a mis-set flag and a deliberate run-down look identical, and only
the owner can tell them apart.
"""
import pytest

from app.portfolio import logic

pytestmark = pytest.mark.regression

CATALOGUE = {
    "B0LIVE1": {"name": "Chana Sattu", "weight": 1.0, "brand": "Mithila Foods", "active": True},
    "B0LIVE2": {"name": "Chana Sattu", "weight": 0.5, "brand": "Mithila Foods", "active": True},
    "B0DEAD1": {"name": "Chana Sattu", "weight": 1.5, "brand": "Mithila Foods", "active": False},
    "B0MOORI": {"name": "Bengali Moori", "weight": 0.5, "brand": "Mithila Foods", "active": False},
    "B0GULAL": {"name": "Herbal Gulal", "weight": 0.5, "brand": "Mithila Foods", "active": False},
}


def _row(child, parent, *, units, sales):
    return {
        "parentAsin": parent,
        "childAsin": child,
        "sales": {
            "orderedProductSales": {"amount": sales, "currencyCode": "INR"},
            "refundedProductSales": {"amount": 0.0, "currencyCode": "INR"},
            "unitsOrdered": units,
            "unitsRefunded": 0,
            "netUnitsSold": units,
        },
        "fees": [],
        "ads": [],
        "netProceeds": {"total": {"amount": sales * 0.3}},
    }


#: A mixed parent (one inactive size WITH sales), a wholly-inactive parent that sold, and one that
#: did not. The inactive size carrying real sales is the part live data cannot supply.
ROWS = [
    _row("B0LIVE1", "P1", units=200, sales=100000.0),
    _row("B0LIVE2", "P1", units=100, sales=30000.0),
    _row("B0DEAD1", "P1", units=50, sales=20000.0),      # inactive size of a LIVE parent
    _row("B0MOORI", "P2", units=121, sales=21633.0),     # wholly inactive, still selling
    _row("B0GULAL", "P3", units=0, sales=0.0),           # wholly inactive, silent
]


def _shown(result):
    return {p["product"] for p in result["parents"]}


# ─── What is hidden, and what the totals then say ─────────────────────────────


def test_an_inactive_size_leaves_its_parents_totals_and_the_account_totals():
    """The filter is at the SIZE level, so a mixed parent keeps its live sizes and loses the rest."""
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    parent = next(p for p in result["parents"] if p["product"] == "Chana Sattu")

    assert {s["asin"] for s in parent["sizes"]} == {"B0LIVE1", "B0LIVE2"}
    assert parent["sales"] == 130000.0, "the inactive size's sales are still in the parent"
    assert result["totals"]["sales"] == 130000.0


def test_a_parents_money_equals_the_sum_of_the_sizes_SHOWN():
    """**The invariant the whole filter has to preserve**, and live data cannot test it.

    All 6 mixed parents on the account have zero sales on their inactive sizes today, so a fixture
    copied from production would pass whether or not the aggregate was scoped. This one puts Rs
    20,000 on the hidden size, where `_sum_sizes(all_sizes)` leaves the parent row claiming more than
    the rows beneath it add up to — the "86 orders beside 87 lines" defect, on a money column.
    """
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    for parent in result["parents"]:
        assert parent["sales"] == pytest.approx(sum(s["sales"] for s in parent["sizes"])), (
            f"{parent['product']} shows {parent['sales']} above rows summing to "
            f"{sum(s['sales'] for s in parent['sizes'])}"
        )
        assert parent["units"] == sum(s["units"] for s in parent["sizes"])


def test_a_wholly_inactive_parent_LEAVES_the_table_and_is_NAMED_with_its_figures():
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    assert _shown(result) == {"Chana Sattu"}
    assert result["inactive_hidden_parents"] == 2
    assert result["inactive_hidden_skus"] == 3, "the mixed parent's hidden size must be counted too"

    named = result["inactive_with_sales"]
    assert [x["product"] for x in named] == ["Bengali Moori"]
    assert named[0]["units"] == 121
    assert named[0]["sales"] == 21633.0


def test_a_wholly_inactive_parent_that_sold_NOTHING_is_counted_but_NOT_named():
    """47 of the 55 hidden parents on the live account sold nothing.

    Naming them would bury the 8 that matter in a list of dead products, which is the opposite of
    what the banner is for. They are still counted, so the row count reconciles.
    """
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    assert "Herbal Gulal" not in {x["product"] for x in result["inactive_with_sales"]}
    assert result["inactive_hidden_parents"] == 2, "the silent parent is still counted"


def test_EVERY_RUPEE_is_either_on_screen_or_named_as_excluded():
    """The property that makes the banner a RECONCILIATION rather than a warning.

    `active_only + inactive_sales == full`, to the paisa — the same assertion
    `test_shipment_catalogue.py` already makes for units, and the reason that tab's banner could be
    trusted. **This is what caught a real bug in my first version**: `inactive_sales` summed the
    vanished PARENTS, so a mixed parent's hidden size was excluded from the totals and counted
    nowhere — Rs 20,000 neither on screen nor named.
    """
    hidden = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    full = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={}, include_inactive=True)

    assert hidden["totals"]["sales"] + hidden["inactive_sales"] == pytest.approx(
        full["totals"]["sales"]
    ), "some sales are neither shown nor reported as excluded"
    assert (
        hidden["totals"]["units_ordered"] + hidden["inactive_sales_units"]
        == full["totals"]["units_ordered"]
    ), "some units are neither shown nor reported as excluded"


def test_include_inactive_restores_every_row():
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={}, include_inactive=True)
    assert _shown(result) == {"Chana Sattu", "Bengali Moori", "Herbal Gulal"}
    assert result["totals"]["sales"] == 171633.0
    assert result["include_inactive"] is True, "the flag must be echoed for the toggle to render"
    parent = next(p for p in result["parents"] if p["product"] == "Chana Sattu")
    assert len(parent["sizes"]) == 3, "the inactive size is back under its parent"


def test_the_named_list_is_CAPPED_while_the_count_and_the_money_stay_exact():
    """"8 products were hidden" is actionable; "some were" is not.

    Live there are exactly 8 such parents, so the "and N more" path never fires on today's data —
    which is why this builds nine. Same discipline as the missing-days list and the catalogue notes.
    """
    catalogue = dict(CATALOGUE)
    rows = []
    for i in range(9):
        asin = f"B0EXTRA{i:03d}"
        catalogue[asin] = {"name": f"Dead Product {i}", "weight": 1.0, "active": False}
        rows.append(_row(asin, f"PX{i}", units=10 + i, sales=1000.0 * (i + 1)))

    result = logic.portfolio(rows, catalogue, ratings={}, decisions={})
    assert len(result["inactive_with_sales"]) == logic.INACTIVE_SHOWN == 8
    assert result["inactive_with_sales_count"] == 9, "the COUNT must not be capped with the list"
    assert result["inactive_hidden_parents"] == 9
    assert result["inactive_sales"] == pytest.approx(sum(1000.0 * (i + 1) for i in range(9)))
    assert result["inactive_sales_units"] == sum(10 + i for i in range(9))


def test_the_named_list_is_ordered_by_SALES_and_not_by_units():
    """This tab is about money; the Shipment tab's equivalent sorts by units because a plan IS units.

    Built so the two orders disagree: Bengali Moori sold MORE units than Moringa Powder and less
    money, which is the real shape on the account (121u/Rs 21,633 against 70u/Rs 31,845).
    """
    catalogue = {
        "B0FEWBIG": {"name": "Moringa Powder", "weight": 0.5, "active": False},
        "B0MANYSM": {"name": "Bengali Moori", "weight": 0.5, "active": False},
    }
    rows = [
        _row("B0FEWBIG", "PA", units=70, sales=31845.0),
        _row("B0MANYSM", "PB", units=121, sales=21633.0),
    ]
    named = logic.portfolio(rows, catalogue, ratings={}, decisions={})["inactive_with_sales"]
    assert [x["product"] for x in named] == ["Moringa Powder", "Bengali Moori"]


# ─── The defaults that keep a row visible ────────────────────────────────────


def test_only_an_EXPLICIT_falsy_flag_hides_a_row():
    """**Three shapes must all stay visible, and none of them occurs in production.**

    `catalogue.is_active` states the rule this follows — "missing data is not a decision" — and each
    branch guards a different disaster:

    * an ASIN absent from the sheet: all 3 decisions stored on production are on such ASINs, so
      reading absence as inactive makes every one of them unreachable at once;
    * an EMPTY catalogue, which is exactly what `load_catalogue` returns during a Google outage:
      reading that as inactive empties the tab from 90 products to 0;
    * an entry present with no `active` key at all. Every `load_catalogue` path sets it, so this
      cannot happen today — and "cannot happen" is how `delete_draft_plans`' docstring came to be
      wrong and destroyed 400 units of packed stock.
    """
    row = _row("B0LIVE1", "P1", units=100, sales=50000.0)
    for label, catalogue in (
        ("explicitly active", {"B0LIVE1": {"name": "X", "weight": 1.0, "active": True}}),
        ("no active key", {"B0LIVE1": {"name": "X", "weight": 1.0}}),
        ("ASIN absent", {"B0OTHER": {"name": "Y", "weight": 1.0, "active": True}}),
        ("empty catalogue (sheet outage)", {}),
    ):
        result = logic.portfolio([row], catalogue, ratings={}, decisions={})
        assert len(result["parents"]) == 1, f"{label}: the row was hidden"
        assert result["inactive_hidden_parents"] == 0, f"{label}: reported as hidden"
        assert result["totals"]["sales"] == 50000.0, f"{label}: the money left the totals"

    hidden = logic.portfolio(
        [row], {"B0LIVE1": {"name": "X", "weight": 1.0, "active": False}},
        ratings={}, decisions={},
    )
    assert not hidden["parents"], "an explicit Active = N did NOT hide the row"
    assert hidden["inactive_sales"] == 50000.0


def test_an_unmatched_ASIN_is_kept_AND_still_reported_as_unmatched():
    """Both halves: keeping it is the Triphala Sattu lesson, reporting it is how the sheet gets fixed.

    A banner naming ASINs that are not in the grid would be a note about rows nobody can see.
    """
    result = logic.portfolio(
        [_row("B0STRANGE", "P9", units=40, sales=9000.0)], CATALOGUE, ratings={}, decisions={},
    )
    assert len(result["parents"]) == 1
    assert result["unmatched_asins"] == ["B0STRANGE"]
    assert result["totals"]["sales"] == 9000.0


# ─── A recorded decision keeps its product visible ───────────────────────────


def test_a_product_with_a_stored_DECISION_is_never_hidden():
    """`ProductDecision` exists to answer "I marked Moori KILL at -56.8% net; what is it now?".

    That question needs the row present — and Bengali Moori is one of the 8 products this filter
    would otherwise remove. Hiding it would leave the table unable to answer the only question the
    decision was recorded for.
    """
    decisions = {"P2": {"decision": "kill", "note": "returns", "decided_at": "2026-08-27"}}
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions=decisions)

    assert "Bengali Moori" in _shown(result)
    moori = next(p for p in result["parents"] if p["product"] == "Bengali Moori")
    assert moori["decision"] == "kill"
    # ...and it SAYS it is inactive, or it renders identically to a live product.
    assert moori["inactive"] is True
    assert result["decided_but_inactive"] == ["Bengali Moori"]
    # Herbal Gulal has no decision and stays hidden.
    assert "Herbal Gulal" not in _shown(result)


def test_a_decided_products_money_is_ON_SCREEN_rather_than_in_the_banner():
    """It is shown, so it must be counted — or the totals and the banner both under-report it."""
    decisions = {"P2": {"decision": "kill", "note": "", "decided_at": "2026-08-27"}}
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions=decisions)
    assert result["totals"]["sales"] == pytest.approx(130000.0 + 21633.0)
    assert "Bengali Moori" not in {x["product"] for x in result["inactive_with_sales"]}


# ─── The verdict must not name a row that is not on screen ───────────────────


def test_a_SURGICAL_reason_never_names_a_hidden_size():
    """Rule 4 walks `sizes` for loss-makers and NAMES them in its reason.

    Given the unfiltered list it can judge a parent SURGICAL because of a pack that is not on screen —
    a verdict the owner cannot check, which is the one thing these reasons exist to prevent. Built so
    the ONLY loss-making size is inactive: the parent must not be SURGICAL at all.
    """
    catalogue = {
        "B0GOOD": {"name": "Chana Sattu", "weight": 1.0, "active": True},
        "B0LOSS": {"name": "Chana Sattu", "weight": 0.25, "active": False},
    }
    rows = [_row("B0GOOD", "P1", units=400, sales=200000.0)]
    losing = _row("B0LOSS", "P1", units=100, sales=10000.0)
    losing["netProceeds"] = {"total": {"amount": -8000.0}}
    rows.append(losing)

    parent = logic.portfolio(rows, catalogue, ratings={}, decisions={})["parents"][0]
    assert parent["verdict"] != logic.VERDICT_SURGICAL, (
        f"judged SURGICAL on a hidden size: {parent['verdict_reason']}"
    )
    assert "250 g" not in parent["verdict_reason"]

    # With the toggle on it IS surgical, which is the same rule reading the same sizes.
    shown = logic.portfolio(rows, catalogue, ratings={}, decisions={},
                            include_inactive=True)["parents"][0]
    assert shown["verdict"] == logic.VERDICT_SURGICAL


def test_the_family_NAME_is_derived_from_the_sizes_shown():
    """A row must describe the rows beneath it.

    A parent holding 3 flavours of which 2 are inactive is, on screen, a single-flavour product —
    labelling it with a family name derived from flavours that are not rendered is the naming version
    of the SURGICAL defect above.
    """
    catalogue = {
        "B0CHEESE": {"name": "Cheese Roasted Chana", "weight": 1.0, "active": True},
        "B0PERI": {"name": "Peri Peri Roasted Chana", "weight": 1.0, "active": False},
        "B0NIMBU": {"name": "Nimbu Roasted Chana", "weight": 1.0, "active": False},
    }
    rows = [
        _row("B0CHEESE", "P1", units=100, sales=50000.0),
        _row("B0PERI", "P1", units=50, sales=20000.0),
        _row("B0NIMBU", "P1", units=50, sales=20000.0),
    ]
    parent = logic.portfolio(rows, catalogue, ratings={}, decisions={})["parents"][0]
    assert parent["product"] == "Cheese Roasted Chana", (
        "the row is named after flavours that are not on screen"
    )
    assert not parent["flavours"], "a flavour heading was kept for hidden flavours"

    grouped = logic.portfolio(rows, catalogue, ratings={}, decisions={},
                              include_inactive=True)["parents"][0]
    assert grouped["product"] == "Roasted Chana", "the family label is wrong with all three shown"


# ─── Sizes hidden from a parent that is still shown ──────────────────────────


def test_sizes_dropped_from_a_SHOWN_parent_are_reported_separately():
    """A different question from a vanished product, so its own line.

    "Should this be selling at all" versus "is this pack size really retired" — and the parent row
    beside the second one looks entirely normal, which makes it the easier one to miss.
    """
    result = logic.portfolio(ROWS, CATALOGUE, ratings={}, decisions={})
    assert result["inactive_sizes_of_shown"] == [{"product": "Chana Sattu", "sizes": 1}]
    # A wholly-hidden parent belongs in the other list, not this one.
    assert "Bengali Moori" not in {x["product"] for x in result["inactive_sizes_of_shown"]}


# ─── Through the real routes ─────────────────────────────────────────────────
#
# `_dashboard` loads the catalogue live, so these patch it: a route test whose answer depends on
# today's MRP sheet would pass or fail on whatever the owner edited this morning.

import json                                                            # noqa: E402
from pathlib import Path                                               # noqa: E402

import pytest_asyncio                                                  # noqa: E402

from app.portfolio import repository                                   # noqa: E402
from app.routers import portfolio as router                            # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
DAY = "2026-08-26"


@pytest_asyncio.fixture
async def seeded(db, monkeypatch):
    """Two products on one day: one active, one inactive and still selling."""
    rows = _stamped([
        _row("B0LIVE1", "P1", units=200, sales=100000.0),
        _row("B0MOORI", "P2", units=121, sales=21633.0),
    ])
    await repository.save_economics_daily(db, rows)

    async def fake_catalogue():
        return CATALOGUE, None, "sheet"

    monkeypatch.setattr(router.catalogue, "load_catalogue", fake_catalogue)
    return rows


def _stamped(rows):
    for row in rows:
        row["startDate"] = DAY
        row["endDate"] = DAY
    return rows


async def test_the_route_hides_inactive_products_and_include_inactive_restores_them(
    auth_client, seeded
):
    hidden = (await auth_client.get(f"/portfolio?start={DAY}&end={DAY}")).json()
    assert [p["product"] for p in hidden["parents"]] == ["Chana Sattu"]
    assert hidden["include_inactive"] is False
    assert hidden["inactive_sales"] == 21633.0
    assert hidden["inactive_sales_units"] == 121

    shown = (
        await auth_client.get(f"/portfolio?start={DAY}&end={DAY}&include_inactive=1")
    ).json()
    assert {p["product"] for p in shown["parents"]} == {"Chana Sattu", "Bengali Moori"}
    assert shown["include_inactive"] is True, "the flag must be echoed so the toggle can render it"
    assert shown["totals"]["sales"] == hidden["totals"]["sales"] + hidden["inactive_sales"]


async def test_a_DECISION_on_a_hidden_product_still_records_its_FIGURES(auth_client, seeded, db):
    """**The non-obvious breakage of hiding rows, and it is silent.**

    `save_decision` looks the parent up by `parent_asin` in `data["parents"]`. On the default view a
    hidden product is not there, so `snapshot` stays None and the decision is stored with no figures
    — defeating the one thing `ProductDecision.snapshot_json` exists for, on exactly the products
    most likely to be marked KILL. So that call passes `include_inactive=True`.
    """
    response = await auth_client.post(
        "/portfolio/decision",
        json={"parent_asin": "P2", "decision": "kill", "note": "inactive and returning"},
    )
    assert response.status_code == 200, response.text

    stored = await repository.load_decisions(db)
    snapshot = stored["P2"]["snapshot"]
    assert snapshot, "a decision on an inactive product recorded no figures at all"
    assert snapshot["sales"] == 21633.0
    assert snapshot["units"] == 121
    assert snapshot["verdict"], "the verdict at the time was not recorded"


async def test_the_WORKBOOK_follows_the_toggle_and_names_what_it_excluded(auth_client, seeded):
    """A file holding different products from the grid it came from is worse than no file.

    The subtitle has to state the exclusion in RUPEES, because a workbook leaves the app without the
    screen's banner beside it — the same reason the pre-COGS caveat is written into row 1.
    """
    import io

    from openpyxl import load_workbook

    def _rows_of(content):
        sheet = load_workbook(io.BytesIO(content)).active
        return list(sheet.iter_rows(values_only=True))

    hidden = _rows_of(
        (await auth_client.get(f"/portfolio/download.xlsx?start={DAY}&end={DAY}")).content
    )
    shown = _rows_of(
        (
            await auth_client.get(
                f"/portfolio/download.xlsx?start={DAY}&end={DAY}&include_inactive=1"
            )
        ).content
    )
    assert len(shown) > len(hidden), "include_inactive did not reach the workbook"

    # `build_portfolio_xlsx` inserts the subtitle ABOVE the header row, so it is row 1 — the first
    # thing read, which is the point of putting the caveats there.
    subtitle = " ".join(str(c) for c in hidden[0] if c)
    assert "Active=N" in subtitle, f"the file does not say it excluded anything: {subtitle!r}"
    assert "21,633" in subtitle, "the excluded RUPEES are not stated, so the total cannot reconcile"
    assert "121 units" in subtitle
    # With the toggle on there is nothing excluded, so the caveat must NOT appear — a caveat that
    # fires on every render is the kind that trains its reader to skip the one that matters.
    assert "Active=N" not in " ".join(str(c) for c in shown[0] if c)
