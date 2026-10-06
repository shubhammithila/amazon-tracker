"""The Portfolio column picker: hide and reorder columns, saved per login.

Asked for as "make every column except the Sales, ad spend, units and weights… hiddenable", then
"also make the columns slidable — which one first, which later, or last". See
docs/superpowers/specs/2026-09-30-portfolio-column-picker-design.md.
"""
import pytest

pytestmark = pytest.mark.regression


def test_users_has_a_nullable_preferences_column():
    """NULL means "never chosen", deliberately distinct from "chose the default": it is what lets a
    later change of default reach everyone who never customised."""
    from app.models import User

    column = User.__table__.c.preferences_json
    assert column.nullable is True
    assert column.server_default is None


from app.portfolio import columns as C

ALL = ["verdict", "sales", "refunded", "fees_total", "ad_spend", "net", "tacos", "acos", "net_pct",
       "units_ordered", "units", "weight_ordered_kg", "weight_kg", "returns_pct", "rating", "decision"]


def test_the_vocabulary_is_exactly_the_agreed_seventeen_columns():
    assert [c["id"] for c in C.COLUMNS] == ["product"] + ALL
    locked = {c["id"] for c in C.COLUMNS if c["locked"]}
    assert locked == {"product", "sales", "ad_spend", "units_ordered", "units", "weight_ordered_kg", "weight_kg"}


def test_the_default_is_todays_screen():
    assert C.DEFAULT_LAYOUT == {"order": ALL, "hidden": ["returns_pct"]}


@pytest.mark.parametrize("junk", [None, "", "[]", 42, [], {"order": "x"}, {"order": [1, 2]}])
def test_anything_malformed_becomes_the_default(junk):
    assert C.normalise_column_layout(junk) == C.DEFAULT_LAYOUT


def test_a_valid_layout_survives_untouched():
    layout = {"order": list(reversed(ALL)), "hidden": ["acos", "rating"]}
    assert C.normalise_column_layout(layout) == layout


def test_unknown_ids_and_duplicates_are_dropped():
    out = C.normalise_column_layout(
        {"order": ["gone", "sales", "sales"] + ALL, "hidden": ["gone", "acos", "acos"]}
    )
    # First occurrence kept: `sales` is the first KNOWN id in the input, so it leads.
    assert out["order"] == ["sales"] + [c for c in ALL if c != "sales"]
    assert out["hidden"] == ["acos"]


def test_product_is_never_stored_in_the_order():
    out = C.normalise_column_layout({"order": ["product"] + ALL, "hidden": []})
    assert "product" not in out["order"]


def test_a_protected_column_cannot_be_hidden():
    out = C.normalise_column_layout(
        {"order": ALL, "hidden": ["sales", "ad_spend", "units", "weight_kg", "product", "acos"]}
    )
    assert out["hidden"] == ["acos"]


def test_a_NEW_column_appears_at_its_default_position_for_an_existing_user():
    """**The rule most likely to be got wrong.** A column added next month is absent from every
    saved layout; dropping it would make it silently never show for anyone who customised. It goes
    after the nearest default-order predecessor the user still has."""
    saved = [c for c in reversed(ALL) if c != "tacos"]      # user reversed everything; no tacos
    out = C.normalise_column_layout({"order": saved, "hidden": []})
    # In the default order tacos follows net (Sales - Refunds - Amazon fees - Ad spend = Net, then
    # TACOS), so it lands right after net here.
    assert out["order"].index("tacos") == out["order"].index("net") + 1
    assert sorted(out["order"]) == sorted(ALL)


def test_a_missing_FIRST_column_goes_to_the_front():
    saved = [c for c in ALL if c != "verdict"]
    out = C.normalise_column_layout({"order": list(reversed(saved)), "hidden": []})
    assert out["order"][0] == "verdict"


def test_the_normaliser_is_idempotent():
    layout = {"order": ["rating", "gone", "sales"], "hidden": ["sales", "decision"]}
    once = C.normalise_column_layout(layout)
    assert C.normalise_column_layout(once) == once


# ── Reading and writing the preference, through the real routes ────────────────────────────

import json

from httpx import ASGITransport, AsyncClient

from app.main import app
from app.routers.auth import SESSION_COOKIE, serializer


async def _named_client(db, username="owner"):
    from app import users as users_repo
    await users_repo.create(db, username=username, full_name=username, is_admin=True,
                            created_by="test")
    transport = ASGITransport(app=app)
    client = AsyncClient(transport=transport, base_url="http://test", follow_redirects=False)
    client.cookies.set(SESSION_COOKIE, serializer.dumps(
        {"authenticated": True, "username": username, "role": "admin"}))
    return client


async def _stored(db, username):
    from sqlalchemy import select
    from app.models import User
    db.expire_all()
    raw = (await db.execute(select(User.preferences_json).where(User.username == username))).scalar()
    return json.loads(raw) if raw else None


async def test_GET_serves_the_vocabulary_and_the_default_for_a_new_login(db):
    async with await _named_client(db) as client:
        body = (await client.get("/portfolio")).json()
    assert [c["id"] for c in body["columns"]] == ["product"] + ALL
    assert body["column_layout"] == C.DEFAULT_LAYOUT
    assert body["column_scope"] == "account"


async def test_a_saved_layout_comes_back_normalised(db):
    async with await _named_client(db) as client:
        # The page always sends the FULL order; a stale id and a protected hide are what to clean.
        order = ["rating"] + [c for c in ALL if c != "rating"] + ["gone"]
        r = await client.put("/portfolio/column-prefs",
                             json={"order": order, "hidden": ["sales", "acos"]})
        assert r.status_code == 200
        saved = r.json()["column_layout"]
        assert saved["order"][0] == "rating" and "gone" not in saved["order"]
        assert saved["hidden"] == ["acos"]
        assert (await client.get("/portfolio")).json()["column_layout"] == saved


async def test_PUT_writes_ONLY_the_session_account_whatever_the_body_says(db):
    from app import users as users_repo
    await users_repo.create(db, username="victim", full_name="v", is_admin=False,
                            created_by="test")
    async with await _named_client(db, "owner") as client:
        await client.put("/portfolio/column-prefs",
                         json={"username": "victim", "order": ALL, "hidden": ["acos"]})
    assert await _stored(db, "victim") is None
    assert (await _stored(db, "owner"))[C.PREFERENCE_KEY]["hidden"] == ["acos"]


async def test_saving_MERGES_so_another_screens_preference_survives(db):
    from sqlalchemy import update
    from app.models import User
    async with await _named_client(db) as client:
        await db.execute(update(User).where(User.username == "owner")
                         .values(preferences_json=json.dumps({"orders_tab": {"x": 1}})))
        await db.commit()
        await client.put("/portfolio/column-prefs", json={"order": ALL, "hidden": []})
    stored = await _stored(db, "owner")
    assert stored["orders_tab"] == {"x": 1}
    assert stored[C.PREFERENCE_KEY] == {"order": ALL, "hidden": []}


async def test_a_shared_password_session_is_told_to_use_the_browser(auth_client):
    body = (await auth_client.get("/portfolio")).json()
    assert body["column_scope"] == "browser"
    assert body["column_layout"] == C.DEFAULT_LAYOUT
    r = await auth_client.put("/portfolio/column-prefs", json={"order": ALL, "hidden": []})
    assert r.status_code == 409


async def test_a_non_JSON_body_is_a_400(db):
    async with await _named_client(db) as client:
        r = await client.put("/portfolio/column-prefs", content=b"not json",
                             headers={"content-type": "application/json"})
    assert r.status_code == 400


async def test_a_corrupt_stored_value_still_renders_the_default(db):
    from sqlalchemy import update
    from app.models import User
    async with await _named_client(db) as client:
        await db.execute(update(User).where(User.username == "owner")
                         .values(preferences_json="{not json"))
        await db.commit()
        assert (await client.get("/portfolio")).json()["column_layout"] == C.DEFAULT_LAYOUT


async def test_the_EXCEL_ignores_the_saved_layout(db):
    """The file leaves the app; a sheet missing TACOS because it was hidden on screen weeks ago is
    a misleading document. Header row identical with and without a layout."""
    async with await _named_client(db) as client:
        before = (await client.get("/portfolio/download.xlsx")).content
        await client.put("/portfolio/column-prefs",
                         json={"order": list(reversed(ALL)), "hidden": ["tacos", "acos"]})
        after = (await client.get("/portfolio/download.xlsx")).content
    assert before == after or _xlsx_rows(before) == _xlsx_rows(after)


def _xlsx_rows(content):
    import io
    from openpyxl import load_workbook
    return [tuple(r) for r in load_workbook(io.BytesIO(content)).active.iter_rows(values_only=True)]
