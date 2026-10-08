"""The Portfolio download: exactly the rows and columns on screen, numbers stored as numbers.

Asked for as *"if I have just selected sattu… download only sattu"*, *"text which are numbers stored
as numbers or %"*, *"same format as what we see on the screen"*, *"no additional commentary"*.
"""
import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.portfolio import columns as column_vocab
from app.portfolio import export

pytestmark = pytest.mark.regression
sys.path.insert(0, str(Path(__file__).parent))
from test_portfolio_api import _seed_snapshot  # noqa: E402

T = Path(__file__).parent.parent / "templates" / "portfolio.html"


async def _screen(auth_client, db):
    await _seed_snapshot(db)
    return (await auth_client.get("/portfolio")).json()


async def _post(auth_client, **body):
    body.setdefault("format", "xlsx")
    return await auth_client.post("/portfolio/export", json=body)


def _sheet(response):
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content)).active


def _grid(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


# ── which rows ──────────────────────────────────────────────────────────────────────────────────

async def test_only_the_rows_the_screen_sent_in_the_screens_order(auth_client, db):
    data = await _screen(auth_client, db)
    parents = data["parents"]
    assert len(parents) >= 2, "the fixture needs two products to prove selection"
    chosen = [parents[1]["parent_asin"], parents[0]["parent_asin"]]          # reversed order
    ws = _sheet(await _post(auth_client, ids=chosen))
    top = [r for r in range(3, ws.max_row + 1) if not ws.row_dimensions[r].outlineLevel]
    assert [ws.cell(r, 4).value for r in top] == chosen
    assert ws.cell(1, 1).value == "Total — 2 product(s)"


async def test_a_filter_that_matches_nothing_downloads_an_empty_table_not_everything(auth_client, db):
    await _screen(auth_client, db)
    ws = _sheet(await _post(auth_client, ids=[]))
    assert ws.max_row == 2 and ws.cell(1, 1).value == "Total — 0 product(s)"


async def test_unknown_ids_are_ignored_and_ids_are_required(auth_client, db):
    data = await _screen(auth_client, db)
    good = data["parents"][0]["parent_asin"]
    ws = _sheet(await _post(auth_client, ids=["B0NOTREAL1", good]))
    assert ws.cell(1, 1).value == "Total — 1 product(s)"
    assert (await _post(auth_client)).status_code == 400
    assert (await _post(auth_client, ids=[1, 2])).status_code == 400
    assert (await _post(auth_client, ids=[], format="csv")).status_code == 400


async def test_the_sku_view_downloads_pack_sizes(auth_client, db):
    data = await _screen(auth_client, db)
    sku = data["skus"][0]["asin"]
    ws = _sheet(await _post(auth_client, ids=[sku], view="skus"))
    assert ws.cell(1, 1).value == "Total — 1 pack size(s)"
    assert ws.cell(3, 4).value == sku and ws.max_row == 3


async def test_sizes_are_grouped_under_their_product_and_open_only_where_open_on_screen(auth_client, db):
    data = await _screen(auth_client, db)
    with_sizes = [p for p in data["parents"] if p["sizes"]][:2]
    ids = [p["parent_asin"] for p in with_sizes]
    ws = _sheet(await _post(auth_client, ids=ids, open=[ids[0]]))
    rows = list(range(3, ws.max_row + 1))
    child = [r for r in rows if ws.row_dimensions[r].outlineLevel]
    assert child, "sizes must be in the file"
    first_parent_children, second = [], []
    current = None
    for r in rows:
        if not ws.row_dimensions[r].outlineLevel:
            current = ws.cell(r, 4).value
        elif current == ids[0]:
            first_parent_children.append(r)
        else:
            second.append(r)
    assert first_parent_children and not any(ws.row_dimensions[r].hidden for r in first_parent_children)
    if len(with_sizes) > 1:
        assert second and all(ws.row_dimensions[r].hidden for r in second)


# ── which columns ───────────────────────────────────────────────────────────────────────────────

async def test_the_columns_are_the_screens_in_the_screens_order(auth_client, db):
    data = await _screen(auth_client, db)
    cols = ["net_pct", "sales", "tacos", "rating"]
    ws = _sheet(await _post(auth_client, ids=[data["parents"][0]["parent_asin"]], columns=cols))
    labels = {c["id"]: c["label"] for c in column_vocab.COLUMNS}
    assert [c.value for c in ws[2]] == ["Product", "Brand", "Size", "ASIN",
                                        labels["net_pct"], labels["sales"], labels["tacos"],
                                        labels["rating"], "Reviews"]


def test_every_screen_column_has_an_export_kind():
    ids = {c["id"] for c in column_vocab.COLUMNS} - {"product"}
    assert ids == set(export.COLUMN_KINDS), "a screen column would vanish from the download"


async def test_an_unknown_column_is_dropped(auth_client, db):
    data = await _screen(auth_client, db)
    ws = _sheet(await _post(auth_client, ids=[data["parents"][0]["parent_asin"]],
                            columns=["sales", "evil"]))
    assert [c.value for c in ws[2]][4:] == [dict((c["id"], c["label"]) for c in column_vocab.COLUMNS)["sales"]]


# ── types and formats ───────────────────────────────────────────────────────────────────────────

async def test_numbers_are_numbers_with_screen_like_formats(auth_client, db):
    data = await _screen(auth_client, db)
    p = next(p for p in data["parents"] if p["sales"] and p["tacos"] is not None)
    cols = ["sales", "tacos", "net_pct", "units", "weight_kg"]
    ws = _sheet(await _post(auth_client, ids=[p["parent_asin"]], columns=cols))
    sales, tacos, net, units, weight = (ws.cell(3, c) for c in range(5, 10))
    assert isinstance(sales.value, (int, float)) and "₹" in sales.number_format
    assert sales.value == pytest.approx(p["sales"])
    assert tacos.value == pytest.approx(p["tacos"]) and tacos.number_format == "0.0%"
    assert net.value == pytest.approx(p["net_pct"]) and net.number_format.startswith("+0.0%")
    assert units.value == p["units"] and isinstance(units.value, int)
    if p["weight_kg"] is not None:
        assert weight.value == pytest.approx(p["weight_kg"]) and '" kg"' in weight.number_format
    import re
    number_as_text = re.compile(r"^[+\-−]?₹?[\d,]+(\.\d+)?\s*(%|kg)?$")
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            assert not (isinstance(cell.value, str) and number_as_text.match(cell.value.strip())), \
                f"{cell.coordinate} holds a number as text: {cell.value!r}"


def test_a_missing_figure_is_a_blank_cell_never_zero():
    row = {"parent_asin": "P", "product": "X", "brand": "B", "sales": 0, "tacos": None,
           "net_pct": None, "acos": None, "weight_kg": None, "returns_pct": 0, "sizes": []}
    table = export.build_table({"parents": [row]},
                               columns=["tacos", "net_pct", "acos", "weight_kg", "returns_pct"])
    values = table.rows[0].values
    assert [values[c] for c in ("tacos", "net_pct", "acos", "weight_kg", "returns_pct")] == [None] * 5


def test_acos_spend_with_no_sales_reads_no_sales():
    table = export.build_table({"parents": [{"parent_asin": "P", "acos_infinite": True,
                                             "sizes": []}]}, columns=["acos"])
    assert table.rows[0].values["acos"] == "no sales"


async def test_no_commentary_row_1_is_totals_row_2_headings(auth_client, db):
    data = await _screen(auth_client, db)
    ws = _sheet(await _post(auth_client, ids=[p["parent_asin"] for p in data["parents"]]))
    assert ws.cell(1, 1).value.startswith("Total — ")
    assert ws.cell(2, 1).value == "Product"
    assert ws.freeze_panes == "E3" and ws.auto_filter.ref.startswith("A2:")
    assert all(c.font.name == "Arial" for c in ws[3])
    assert all(c.font.name == "Arial" and c.font.bold for c in ws[2]), "headings must be bold Arial"
    assert all(c.font.bold for c in ws[1]), "the totals row must stand out"
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert len(cell.value) <= 60, f"prose in {cell.coordinate}: {cell.value[:60]!r}"


# ── totals ──────────────────────────────────────────────────────────────────────────────────────

async def test_totals_match_the_pages_own_computeTotals(auth_client, db):
    """The Python port is pinned to the page's JavaScript, run on the same rows."""
    import json

    from tests.js_harness import run_portfolio_js

    data = await _screen(auth_client, db)
    rows = data["parents"]
    js = run_portfolio_js(f"""
const t = computeTotals({json.dumps(rows)});
emit({{sales: t.sales, spend: t.spend, tacos: t.ratio(t.spend, t.sales),
      net: t.ratio(t.net, t.sales), refunds: t.ratio(t.refundedSales, t.sales),
      fees: t.ratio(t.fees, t.sales), units: t.units, ordered: t.ordered,
      weight: t.weighed ? t.weight : null, rating: t.rating, reviews: t.reviews,
      returns: t.ratio(t.refunded, t.ordered)}});""")
    py = export.totals(rows)
    pairs = [("sales", "sales"), ("ad_spend", "spend"), ("tacos", "tacos"), ("net_pct", "net"),
             ("refunds_pct", "refunds"), ("fees_pct", "fees"), ("units", "units"),
             ("units_ordered", "ordered"), ("weight_kg", "weight"), ("rating", "rating"),
             ("returns_pct", "returns")]
    for ours, theirs in pairs:
        if js[theirs] is None:
            assert py[ours] is None, ours
        else:
            assert py[ours] == pytest.approx(js[theirs]), ours
    if py["rating"] is not None:
        assert py["reviews"] == js["reviews"]


def test_parent_percentages_are_recomputed_from_sums_never_averaged():
    rows = [{"sales": 9000, "ad_spend": 900}, {"sales": 1000, "ad_spend": 500}]
    assert export.totals(rows)["tacos"] == pytest.approx(0.14)        # not (10% + 50%) / 2


# ── PDF ─────────────────────────────────────────────────────────────────────────────────────────

async def test_the_pdf_has_the_visible_rows_and_the_rupee_font(auth_client, db):
    from pypdf import PdfReader

    data = await _screen(auth_client, db)
    p = data["parents"][0]
    r = await _post(auth_client, ids=[p["parent_asin"]], format="pdf", columns=["sales"],
                    label="Sattu")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert "sattu" in r.headers["content-disposition"]
    text = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
    assert p["product"].split()[0] in text and "₹" in text and "Sattu" in text
    # Collapsed sizes are left out, as on screen.
    for size in p["sizes"]:
        assert size["asin"] not in text


# ── the page ────────────────────────────────────────────────────────────────────────────────────

def test_the_page_sends_the_visible_rows_columns_and_open_state():
    src = T.read_text(encoding="utf-8")
    body = src[src.index("function exportPayload("):src.index("async function downloadExport(")]
    assert "visible().map(" in body, "the rows must be the ones visible() returns"
    assert "visibleColumns().slice(1)" in body
    assert "[...open]" in body and "filterLabel()" in body
    assert 'data-export="xlsx"' in src and 'data-export="pdf"' in src
    assert "excel-link" not in src


def test_the_filename_names_the_window_and_the_filter():
    assert export.filename(("2026-09-08", "2026-10-07"), 'Sattu · search "jau"', "xlsx") == \
        "portfolio-2026-09-08_2026-10-07-sattu-search-jau.xlsx"
    assert export.filename(None, "", "pdf") == "portfolio.pdf"
