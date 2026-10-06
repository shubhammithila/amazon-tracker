"""Refunds, Amazon fees and Net (₹) beside Sales and Ad spend: a row reads as Amazon's own sum.

Asked for as *"add a column showing amazon fees combined except for the ads cost… for each parent
and the child skus… I am selling at 100 - 40 (amazon fees) - 30 (ads) = net 30"*. Amazon's net also
takes REFUNDS off, so a fees column alone would leave a row that does not add up on screen — the
property these columns exist to make visible, and the one asserted here on every row.
"""
import io
import sys
from pathlib import Path

import pytest

from tests.js_harness import run_portfolio_js

pytestmark = pytest.mark.regression

sys.path.insert(0, str(Path(__file__).parent))
from test_portfolio_columns_render import VOCAB  # noqa: E402


async def test_EVERY_row_adds_up_sales_minus_refunds_minus_fees_minus_ads_equals_net(auth_client, db):
    """On every size AND parent row of the real-shaped fixture. A parent is checked separately from
    its sizes because it is built by `_sum_sizes`, which is where a missing `refunded` lived."""
    from test_portfolio_api import _seed_snapshot

    await _seed_snapshot(db)
    data = (await auth_client.get("/portfolio")).json()
    rows = [(p["product"], p) for p in data["parents"]]
    rows += [(f"{p['product']} / {s['asin']}", s) for p in data["parents"] for s in p["sizes"]]
    assert rows
    for label, r in rows:
        assert "refunded" in r and "fees_total" in r, f"{label} carries no refunds or fees figure"
        gap = r["sales"] - r["refunded"] - r["fees_total"] - r["ad_spend"] - r["net"]
        assert abs(gap) < 0.05, f"{label}: the row does not add up, off by ₹{gap:.2f}"
    assert any(r["refunded"] for _, r in rows), "the fixture has no refunds, so the identity proves less"


def test_the_parent_sums_its_sizes_refunds():
    from app.portfolio import logic

    def size(asin, refunded):
        return logic.size_row(
            {"childAsin": asin, "parentAsin": "B0P",
             "sales": {"orderedProductSales": {"amount": 1000}, "refundedProductSales": {"amount": refunded},
                       "unitsOrdered": 10, "netUnitsSold": 9}},
            {asin: {"name": "X", "weight": 1}},
        )

    parent = logic._sum_sizes([size("B0A", 120.0), size("B0B", 30.5)])
    assert parent["refunded"] == 150.5


@pytest.mark.parametrize("col, expected", [
    ("refunded", "₹218"), ("fees_total", "₹1,461"), ("net", "₹914"),
])
def test_each_money_column_shows_its_own_figure(col, expected):
    """Values all DIFFERENT, so a renderer reading a neighbour's field cannot pass."""
    out = run_portfolio_js(VOCAB + """
const R = Object.assign({}, ROW, {refunded: 217.6, fees_total: 1461.2, net: 913.7, ad_spend: 1640});
layout = normaliseLayout(null, data.columns);
const part = dataCells(R).split('data-col="%s"')[1] || "";
emit(part ? part.slice(part.indexOf(">") + 1, part.indexOf("</td>")).replace(/<[^>]+>/g, "").trim() : null);
""" % col)
    assert out == expected


def test_a_LOSS_is_red_with_a_real_minus_sign():
    out = run_portfolio_js("emit(netMoney(-5400.4))")
    assert out == '<span class="neg">−₹5,400</span>'
    assert run_portfolio_js("emit(netMoney(0))") == "₹0"


def test_the_totals_row_sums_fees_and_refunds_separately():
    out = run_portfolio_js(VOCAB + """
const A = Object.assign({}, ROW, {refunded: 100, fees_total: 1000});
const B = Object.assign({}, ROW, {refunded: 50, fees_total: 300});
layout = normaliseLayout(null, data.columns);
const t = computeTotals([A, B]);
emit([t.refundedSales, t.fees]);
""")
    assert out == [150, 1300]


def test_the_columns_read_in_the_order_of_the_sum():
    from app.portfolio import columns as C

    order = [c["id"] for c in C.COLUMNS]
    assert order.index("sales") < order.index("refunded") < order.index("fees_total") \
        < order.index("ad_spend") < order.index("net")
    labels = {c["id"]: c["label"] for c in C.COLUMNS}
    assert labels["fees_total"] == "Amazon fees" and labels["net"] == "Net (₹)"


def test_every_new_column_sorts_by_itself_not_by_sales():
    out = run_portfolio_js("""
emit(["refunded", "fees_total", "net"].filter(id => !FIELDS[COLUMN_DEFS[id].sortKey]));
""")
    assert out == []


async def test_the_workbook_carries_refunds_and_fees_matching_the_screen(auth_client, db):
    from openpyxl import load_workbook

    from test_portfolio_api import _seed_snapshot

    await _seed_snapshot(db)
    screen = (await auth_client.get("/portfolio")).json()
    sheet = load_workbook(io.BytesIO((await auth_client.get("/portfolio/download.xlsx")).content)).active
    rows = [[c.value for c in r] for r in sheet.iter_rows()]
    at = next(i for i, r in enumerate(rows) if r and r[0] == "Product")
    h = rows[at]
    asin, ref, fee, net = h.index("ASIN"), h.index("Refunds"), h.index("Amazon fees"), h.index("Net")
    by_asin = {r[asin]: (r[ref], r[fee], r[net]) for r in rows[at + 1:] if r[asin]}
    for s in (s for p in screen["parents"] for s in p["sizes"]):
        assert by_asin[s["asin"]] == (s["refunded"], s["fees_total"], s["net"]), s["asin"]
