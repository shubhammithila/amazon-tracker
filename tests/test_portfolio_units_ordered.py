"""Units ordered (Seller Central's figure) and net units (after refunds) are BOTH shown, labelled.

Reported as *"there is a difference between the units actually sold as per the business report and
what is being shown"*. For 5 Sep - 4 Oct, Bengali Posta read 106 + 84 on screen against the Business
Report's 113 + 86. The stored data was right: `units_ordered` was 113 and 86, exactly the report. The
column labelled "Units" was `netUnitsSold`, the units ordered MINUS the 7 + 2 refunded, a figure the
Business Report never shows, so the two could never be reconciled by eye.

Measured across the account for the same window: 87 of 114 ASINs matched the report exactly on
`units_ordered`, and the remaining 50 units of 14,785 (0.3%) are replacement orders and pending
lines, which the Business Report counts and the profit API does not.
"""
import io
import sys
from pathlib import Path

import pytest

from tests.js_harness import run_portfolio_js

pytestmark = pytest.mark.regression

sys.path.insert(0, str(Path(__file__).parent))
from test_portfolio_columns_render import VOCAB  # noqa: E402

POSTA = """
const POSTA = Object.assign({}, ROW, {units_ordered: 113, units: 106, units_refunded: 7});
"""


def _cell(html_expr: str, col: str) -> str:
    return run_portfolio_js(VOCAB + POSTA + f"""
layout = normaliseLayout(null, data.columns);
const html = {html_expr};
const m = new RegExp('<td data-col="{col}"[^>]*>([\\\\s\\\\S]*?)</td>').exec(html);
emit(m ? m[1].replace(/<[^>]+>/g, "").trim() : null);
""")


@pytest.mark.parametrize("col, expected", [("units_ordered", "113"), ("units", "106")])
def test_each_unit_column_shows_its_OWN_figure_not_the_other_one(col, expected):
    """Built with DIFFERENT values for the two fields. With equal ones (the shared fixture has
    10 and 10) a renderer reading the wrong field would pass."""
    assert _cell('`<td data-col="product"></td>` + dataCells(POSTA)', col) == expected
    assert _cell("sizeRowHtml(POSTA, false)", col) == expected


def test_the_totals_row_sums_units_ordered_separately():
    out = _cell("totalsRow([POSTA, POSTA], false)", "units_ordered")
    assert out == "226"
    assert _cell("totalsRow([POSTA, POSTA], false)", "units") == "212"


def test_the_two_columns_are_labelled_so_they_cannot_be_confused():
    from app.portfolio import columns as C

    labels = {c["id"]: c["label"] for c in C.COLUMNS}
    assert labels["units_ordered"] == "Units ordered"
    assert labels["units"] == "Net units", "the after-refunds figure must not be called plain 'Units'"
    assert {"units_ordered", "units"} <= {c["id"] for c in C.COLUMNS if c["locked"]}


async def test_the_workbook_carries_both_unit_columns_with_the_right_figures(auth_client, db):
    from openpyxl import load_workbook

    from test_portfolio_api import _seed_snapshot

    await _seed_snapshot(db)
    response = await auth_client.get("/portfolio/download.xlsx")
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = [[c.value for c in r] for r in sheet.iter_rows()]
    header_at = next(i for i, r in enumerate(rows) if r and r[0] == "Product")
    headers = rows[header_at]
    assert "Units ordered" in headers and "Net units" in headers and "Units" not in headers
    ordered, net, asin_col = headers.index("Units ordered"), headers.index("Net units"), headers.index("ASIN")
    by_asin = {r[asin_col]: (r[ordered], r[net]) for r in rows[header_at + 1:] if r[asin_col]}

    # Each size line must carry the SCREEN's two figures for that ASIN. "ordered >= net" alone let a
    # mutation writing net units into BOTH columns pass, since equal values satisfy it.
    screen = (await auth_client.get("/portfolio")).json()
    sizes = [s for p in screen["parents"] for s in p["sizes"]]
    differing = [s for s in sizes if s["units_ordered"] != s["units"]]
    assert differing, "the fixture has no refunded units, so this test could not tell the columns apart"
    for s in sizes:
        assert by_asin[s["asin"]] == (s["units_ordered"], s["units"]), s["asin"]


# ─── Sales ex-GST, and weight on units ORDERED ───────────────────────────────


def test_sales_is_labelled_ex_GST_everywhere_it_is_shown():
    """Amazon's `orderedProductSales` excludes the 5% GST the order price includes — measured,
    app sales = order price / 1.05 to within 0.2% — so the Business Report reads ~5% higher. An
    unlabelled "Sales" invites exactly the reconciliation this was asked about."""
    from pathlib import Path

    from app.portfolio import columns as C

    assert {c["id"]: c["label"] for c in C.COLUMNS}["sales"] == "Sales (ex-GST)"
    text = Path("templates/portfolio.html").read_text(encoding="utf-8")
    assert '<div class="k">Sales (ex-GST)</div>' in text, "the KPI tile still says plain Sales"
    router = Path("app/routers/portfolio.py").read_text(encoding="utf-8")
    assert '"Sales (ex-GST)"' in router, "the Excel header still says plain Sales"


def test_weight_ordered_is_units_ORDERED_times_the_pack_and_sums_up_to_the_parent():
    from app.portfolio import logic

    size = logic.size_row(
        {"childAsin": "B0X", "parentAsin": "B0P",
         "sales": {"orderedProductSales": {"amount": 100}, "unitsOrdered": 113, "unitsRefunded": 7,
                   "netUnitsSold": 106}},
        {"B0X": {"name": "Bengali Posta", "weight": 0.2}},
    )
    assert size["weight_ordered_kg"] == 22.6       # 113 x 0.2, not 106 x 0.2
    assert size["weight_kg"] == 21.2
    parent = logic._sum_sizes([size, dict(size)])
    assert parent["weight_ordered_kg"] == 45.2


def test_an_unknown_pack_weight_is_a_dash_not_zero_for_weight_ordered_too():
    from app.portfolio import logic

    size = logic.size_row(
        {"childAsin": "B0X", "sales": {"unitsOrdered": 5, "netUnitsSold": 5}}, {"B0X": {"name": "X"}}
    )
    assert size["weight_ordered_kg"] is None


@pytest.mark.parametrize("col", ["weight_ordered_kg", "weight_kg"])
def test_each_weight_column_shows_its_own_figure(col):
    expected = {"weight_ordered_kg": "22.6 kg", "weight_kg": "21.2 kg"}[col]
    out = run_portfolio_js(VOCAB + """
const R = Object.assign({}, ROW, {weight_ordered_kg: 22.6, weight_kg: 21.2});
layout = normaliseLayout(null, data.columns);
const part = dataCells(R).split('data-col="%s"')[1] || "";
emit(part ? part.slice(part.indexOf(">") + 1, part.indexOf("</td>")).replace(/<[^>]+>/g, "").trim() : null);
""" % col)
    assert out == expected


def test_EVERY_sortable_column_has_a_sort_field_so_it_does_not_silently_sort_by_sales():
    """**Found while adding Weight ordered.** `sortValue` reads `FIELDS[sort.key] || FIELDS.sales`,
    so a column whose sortKey is missing from FIELDS sorts by SALES with nothing failing — which is
    what "Units ordered" did from the commit that added it. Asserted against the page's own objects."""
    out = run_portfolio_js("""
emit(Object.entries(COLUMN_DEFS).filter(([id, d]) => d.sortKey && !FIELDS[d.sortKey]).map(([id]) => id));
""")
    assert out == [], f"these columns sort by Sales instead of themselves: {out}"
