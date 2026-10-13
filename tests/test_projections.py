"""Portfolio -> Projections: how much stock to hold, live from the Profit tab's sales.

The owner's formula, 12 Oct 2026: daily = 60% x 7-day rate + 40% x 30-day rate; projection =
daily x 30 x seasonal x (1 + growth); Ideal FBA = daily' x lead total x markup; WH buffer = the
slips added up; Ideal WH = daily' x buffer x markup; Total = FBA + WH.
"""
import ast
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy import func, select

from app.models import ProjectionRow
from app.projections import logic, service

pytestmark = pytest.mark.regression
ROOT = Path(__file__).resolve().parent.parent
S = logic.DEFAULT_SETTINGS


# ── the formula ─────────────────────────────────────────────────────────────────────────────────

def chana(**kw):
    """Chana Sattu on production, 11 Oct: 1,468.5 kg in 30 days, 340.48 kg in 7, lead 5/2/10/5."""
    return {"kg_30d": 1468.5, "kg_7d": 340.48, "seasonal_impact": 1.0, "supplier_to_wh": 5,
            "packing": 2, "wh_to_ixd": 10, "ixd_to_fba": 5, **kw}


def test_the_owners_worked_example():
    out = logic.compute(chana(), S)
    assert out["daily_rate"] == pytest.approx(0.6 * 340.48 / 7 + 0.4 * 1468.5 / 30, abs=0.01)
    daily = out["daily_rate"] * 1.3
    assert out["lead_total"] == 22
    assert out["buffer_days"] == pytest.approx(5 * .4 + 2 * .3 + 10 * .5 + 5 * .3)       # 9.1
    assert out["ideal_fba"] == pytest.approx(daily * 22 * 1.5, abs=0.2)
    assert out["ideal_wh"] == pytest.approx(daily * 9.1 * 1.5, abs=0.2)
    assert out["total_stock"] == pytest.approx(out["ideal_fba"] + out["ideal_wh"], abs=0.1)
    assert out["projection_30d"] == pytest.approx(daily * 30, abs=0.2)


def test_the_7_day_figure_weighs_MORE_than_the_30_day_one():
    spike = logic.compute(chana(kg_7d=700.0), S)["daily_rate"]
    assert spike == pytest.approx(0.6 * 100 + 0.4 * 48.95, abs=0.01)
    assert S["seven_day_weight"] > 0.5


def test_ideal_fba_covers_the_FULL_lead_time_including_supplier_to_warehouse():
    a = logic.compute(chana(supplier_to_wh=5), S)
    b = logic.compute(chana(supplier_to_wh=15), S)
    assert b["lead_total"] - a["lead_total"] == 10
    assert b["ideal_fba"] > a["ideal_fba"]


def test_the_buffer_is_the_slips_ADDED_and_the_settings_move_it():
    halved = {**S, "slip_wh_to_ixd": 25.0}
    assert logic.compute(chana(), halved)["buffer_days"] == pytest.approx(9.1 - 2.5)


def test_the_markup_scales_both_halves_and_growth_and_seasonal_scale_the_projection():
    one = logic.compute(chana(), {**S, "stock_markup": 1.0})
    assert logic.compute(chana(), S)["total_stock"] == pytest.approx(one["total_stock"] * 1.5, abs=0.2)
    flat = logic.compute(chana(), {**S, "global_growth_rate": 0.0})
    festive = logic.compute(chana(seasonal_impact=2.0), {**S, "global_growth_rate": 0.0})
    assert festive["projection_30d"] == pytest.approx(flat["projection_30d"] * 2, abs=0.2)


def test_no_7_day_window_is_the_30_day_rate_alone_but_a_quiet_week_is_blended():
    assert logic.daily_rate(1500.0, None, 0.6) == pytest.approx(50.0)
    assert logic.daily_rate(1500.0, 0.0, 0.6) == pytest.approx(20.0)


def test_an_unknown_pack_weight_is_a_dash_never_zero_kg():
    out = logic.compute(chana(kg_30d=None, kg_7d=None), S)
    assert out["total_stock"] is None and out["ideal_fba"] is None and out["lead_total"] == 22
    assert logic.sum_rows([{"total_stock": None}])["total_stock"] is None


def test_settings_are_checked_on_read_and_old_keys_are_ignored():
    got = logic.settings_or_default({"stock_markup": 0.5, "seven_day_weight": 0.7,
                                     "divergence_pct": 30, "global_growth_rate": "x"})
    assert got["stock_markup"] == 1.5 and got["seven_day_weight"] == 0.7
    assert got["global_growth_rate"] == 0.3 and "divergence_pct" not in got


@pytest.mark.parametrize("field,value", [("packing", 2.5), ("packing", -1), ("seasonal_impact", 11),
                                         ("last_month_sale", 100), ("wh_to_ixd", "x")])
def test_edits_outside_the_rules_are_refused(field, value):
    assert logic.edit_error(field, value)


def test_only_seasonal_and_the_four_lead_times_are_editable():
    assert logic.EDITABLE == ("seasonal_impact", "supplier_to_wh", "packing", "wh_to_ixd", "ixd_to_fba")


# ── the live table, from the Profit tab's own rows ──────────────────────────────────────────────

def parent(key, name, kg, **kw):
    return {"parent_asin": key, "product": name, "brand": kw.get("brand", "Mithila Foods"),
            "weight_ordered_kg": kg, "inactive": kw.get("inactive", False), "sizes": [{"product": name}]}


@pytest.fixture
def profit(monkeypatch):
    """Stand in for the Profit tab's `_dashboard`: one answer per window."""
    service._CACHE.clear()
    state = {"30": [parent("B0CHANA", "Chana Sattu", 1468.5), parent("B0JAU", "Jau Sattu", 822.5),
                    parent("B0RICE", "Usna Chawal", 974.0), parent("B0NEW", "Brand New Thing", 50.0),
                    parent("B0NOWT", "Mystery Pack", None)],
             "7": [parent("B0CHANA", "Chana Sattu", 340.48), parent("B0RICE", "Usna Chawal", 231.0)],
             "complete_7": True, "calls": []}

    async def fake(db, window, *, include_inactive=False):
        days = "7" if window == service.windows(service.ist.today())["7"] else "30"
        state["calls"].append((days, include_inactive))
        complete = days == "30" or state["complete_7"]
        return {"parents": state[days], "completeness": {"complete": complete, "missing_count": 0},
                "inactive_hidden_parents": 3}

    async def cats(db):
        return [SimpleNamespace(product_key="chana sattu", priority=1),
                SimpleNamespace(product_key="jau sattu", priority=1),
                SimpleNamespace(product_key="usna chawal", priority=4)]

    monkeypatch.setattr("app.routers.portfolio._dashboard", fake)
    monkeypatch.setattr("app.shipment.repository.load_categories", cats)
    monkeypatch.setattr(service, "_defaults_file", lambda: {
        "Jau Sattu": {"supplier_to_wh": 2, "packing": 2, "wh_to_ixd": 10, "ixd_to_fba": 5,
                      "seasonal_impact": 1.5}})
    yield state
    service._CACHE.clear()


async def test_rows_follow_the_shipment_tabs_category_order_then_biggest_seller(profit, db):
    data = await service.build(db)
    assert [r["product"] for r in data["rows"]] == [
        "Chana Sattu", "Jau Sattu", "Usna Chawal", "Brand New Thing", "Mystery Pack"]
    assert [c["category"] for c in data["categories"]] == ["Sattu", "Rice", "Unclassified"]


async def test_a_product_absent_from_a_COMPLETE_7_day_window_sold_nothing_this_week(profit, db):
    rows = {r["product"]: r for r in (await service.build(db))["rows"]}
    assert rows["Jau Sattu"]["kg_7d"] == 0.0
    profit["complete_7"] = False
    service._CACHE.clear()
    rows = {r["product"]: r for r in (await service.build(db))["rows"]}
    assert rows["Jau Sattu"]["kg_7d"] is None and rows["Chana Sattu"]["kg_7d"] is None


async def test_lead_times_come_from_saved_then_the_defaults_file_then_global(profit, db):
    db.add(ProjectionRow(product_key=None, parent_product="usna chawal", supplier_to_wh=2, packing=2,
                         wh_to_ixd=10, ixd_to_fba=5, seasonal_impact=1.0, sales_source="sheet"))
    await db.commit()
    rows = {r["product"]: r for r in (await service.build(db))["rows"]}
    assert rows["Usna Chawal"]["config_source"] == "saved" and rows["Usna Chawal"]["supplier_to_wh"] == 2
    assert rows["Jau Sattu"]["config_source"] == "defaults" and rows["Jau Sattu"]["seasonal_impact"] == 1.5
    assert rows["Brand New Thing"]["config_source"] == "global"
    assert rows["Brand New Thing"]["supplier_to_wh"] == logic.GLOBAL_DEFAULTS["supplier_to_wh"]


async def test_a_page_load_never_writes(profit, db):
    await service.build(db)
    assert (await db.execute(select(func.count()).select_from(ProjectionRow))).scalar() == 0


async def test_the_inactive_toggle_is_the_profit_tabs_own(profit, db):
    await service.build(db, include_inactive=True)
    assert set(profit["calls"]) == {("30", True), ("7", True)}


async def test_a_category_card_is_the_sum_of_its_rows(profit, db):
    data = await service.build(db)
    sattu = next(c for c in data["categories"] if c["category"] == "Sattu")
    rows = [r for r in data["rows"] if r["category"] == "Sattu"]
    assert sattu["total_stock"] == pytest.approx(sum(r["total_stock"] for r in rows), abs=0.2)
    assert sattu["products"] == 2


# ── the routes ──────────────────────────────────────────────────────────────────────────────────

async def test_saving_one_lead_time_keeps_the_others_it_was_shown(profit, auth_client, read_committed):
    r = await auth_client.post("/projections/save", json={"changes": [{"key": "B0JAU", "packing": 4}]})
    assert r.status_code == 200
    jau = next(x for x in r.json()["rows"] if x["key"] == "B0JAU")
    assert jau["packing"] == 4 and jau["supplier_to_wh"] == 2 and jau["seasonal_impact"] == 1.5
    assert jau["config_source"] == "saved"


async def test_a_bad_cell_refuses_the_WHOLE_save(profit, auth_client, db):
    r = await auth_client.post("/projections/save", json={"changes": [
        {"key": "B0JAU", "packing": 4}, {"key": "B0CHANA", "wh_to_ixd": 2.5}]})
    assert r.status_code == 400 and "Chana Sattu" in r.json()["error"]
    assert (await db.execute(select(func.count()).select_from(ProjectionRow))).scalar() == 0


@pytest.mark.parametrize("change", [{"key": "B0NOPE", "packing": 3},
                                    {"key": "B0JAU", "last_month_sale": 999}])
async def test_unknown_products_and_read_only_fields_are_refused(profit, auth_client, change):
    assert (await auth_client.post("/projections/save", json={"changes": [change]})).status_code == 400


async def test_a_pre_key_row_is_ADOPTED_not_duplicated(profit, auth_client, db):
    db.add(ProjectionRow(parent_product="Usna Chawal", supplier_to_wh=2, packing=2, wh_to_ixd=10,
                         ixd_to_fba=5, seasonal_impact=1.0, sales_source="sheet"))
    await db.commit()
    r = await auth_client.post("/projections/save", json={"changes": [{"key": "B0RICE", "ixd_to_fba": 7}]})
    assert r.status_code == 200
    rows = (await db.execute(select(ProjectionRow))).scalars().all()
    assert len(rows) == 1 and rows[0].product_key == "B0RICE" and rows[0].supplier_to_wh == 2


async def test_settings_save_refuse_and_reset(profit, auth_client):
    assert (await auth_client.post("/projections/settings", json={"stock_markup": 0.8})).status_code == 400
    r = await auth_client.post("/projections/settings", json={"stock_markup": 2.0})
    assert r.json()["settings"]["stock_markup"] == 2.0
    data = (await auth_client.get("/projections/last")).json()
    assert data["settings"]["stock_markup"] == 2.0
    r = await auth_client.post("/projections/settings", json={"reset": True})
    assert r.json()["settings"] == logic.DEFAULT_SETTINGS


@pytest.mark.parametrize("path", ["/projections/download/reorder.xlsx", "/projections/download/reorder.pdf",
                                  "/projections/download/table.xlsx", "/projections/download/table.pdf"])
async def test_the_downloads_carry_total_stock_to_hold(profit, auth_client, path):
    r = await auth_client.get(path)
    assert r.status_code == 200 and len(r.content) > 500
    if path.endswith("xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        headers = [c.value for c in ws[2]]
        assert "Total stock to hold" in headers and "Ideal FBA" in headers
        assert "In FBA / transit" in headers and "To purchase" in headers


async def test_a_login_without_projections_cannot_reach_it(ops_client):
    assert (await ops_client.get("/projections/last")).status_code in (303, 403)


async def test_the_profit_tab_shows_the_SAME_total_stock(profit, auth_client, db):
    """The Profit column is this tab's figure for the same row id, inactive products included."""
    shown = (await auth_client.get("/projections/last?include_inactive=true")).json()["rows"]
    totals = {r["key"]: r["total_stock"] for r in shown}
    from app.routers import portfolio as pf
    data = {"parents": [{"parent_asin": "B0CHANA"}, {"parent_asin": "B0GONE"}]}
    await pf._attach_stock(db, data)
    assert data["parents"][0]["total_stock"] == totals["B0CHANA"] is not None
    assert data["parents"][1]["total_stock"] is None


async def test_a_projections_failure_never_takes_the_profit_tab_down(monkeypatch, db):
    async def boom(db, **kw):
        raise RuntimeError("sheet down")
    monkeypatch.setattr(service, "stock_by_key", boom)
    from app.routers import portfolio as pf
    data = {"parents": [{"parent_asin": "B0CHANA"}]}
    await pf._attach_stock(db, data)
    assert data["parents"][0]["total_stock"] is None


@pytest.mark.parametrize("module", ["app/projections/service.py", "app/projections/logic.py",
                                    "app/projections/repository.py", "app/routers/projections.py"])
def test_projections_can_neither_call_amazon_nor_write_sales(module):
    """It READS the Profit tab's stored days. A second writer of that cache is what overwrote a month
    of fees on 4 Oct."""
    text = (ROOT / module).read_text(encoding="utf-8")
    tree = ast.parse(text)
    called = {n.attr for n in ast.walk(tree) if isinstance(n, ast.Attribute)}
    assert not called & {"fetch_economics", "save_economics_daily", "run_incremental"}
    assert "UploadFile" not in text and "upload-csv" not in text


# ── the page ────────────────────────────────────────────────────────────────────────────────────

PAGE = (ROOT / "templates" / "projections.html").read_text(encoding="utf-8")


def test_save_starts_DISABLED_and_wakes_only_on_a_real_change():
    assert '<button class="btn" id="save-btn" type="button" disabled>' in PAGE
    assert '$("save-btn").disabled = !count || bad;' in PAGE
    assert "if(bad || same) delete mine[field]; else mine[field] = Number(raw);" in PAGE


def test_the_manual_upload_and_remove_controls_are_gone():
    for gone in ("upload", "csv-file", "Reset all manual edits", "Remove selected", "last_month_sale",
                 "Refresh sales now"):
        assert gone not in PAGE, gone


def test_only_the_five_inputs_are_editable_on_screen():
    assert 'const EDITABLE = ["seasonal_impact"].concat(LEGS.map(l => l[0]));' in PAGE
    assert PAGE.count("editCell(") == 2      # its definition and the one call


def test_it_is_a_portfolio_sub_tab():
    assert '{% include "_portfolio_tabs.html" %}' in PAGE
    tabs = (ROOT / "templates" / "_portfolio_tabs.html").read_text(encoding="utf-8")
    assert 'href="/portfolio-page/projections"' in tabs


def test_sorting_by_category_restores_the_packing_order():
    from tests.js_harness import run_template_js
    out = run_template_js(ROOT / "templates" / "projections.html", ["shownRows", "sortedRows"], [], """
$ = () => ({value: ""});
let category = "";
let sort = {key: "category", dir: 1};
data = {categories: [{category: "Sattu"}, {category: "Rice"}],
        rows: [{product: "Usna", category: "Rice", kg_30d: 900}, {product: "Jau", category: "Sattu", kg_30d: 800},
               {product: "Chana", category: "Sattu", kg_30d: 1400}]};
emit(sortedRows().map(r => r.product));
""")
    assert out == ["Chana", "Jau", "Usna"]


def test_an_edit_box_refuses_half_days_and_negatives():
    from tests.js_harness import run_template_js
    out = run_template_js(ROOT / "templates" / "projections.html", ["editBad"], [], """
data = {edit_ranges: {packing: [0, 365], seasonal_impact: [0, 10]}};
emit([editBad("packing", "2.5"), editBad("packing", "-1"), editBad("packing", ""), editBad("packing", "3"),
      editBad("seasonal_impact", "1.25")]);
""")
    assert out == [True, True, True, False, False]


async def test_an_inactive_product_never_shows_until_asked_even_with_a_decision(profit, db):
    """The Profit tab keeps an inactive product that carries a decision (Bengali Chana Sattu); stock
    is never planned for one, so Projections hides every inactive row behind the toggle."""
    profit["30"].append(parent("B0OLD", "Bengali Chana Sattu", 0.0, inactive=True))
    service._CACHE.clear()
    data = await service.build(db)
    assert "Bengali Chana Sattu" not in [r["product"] for r in data["rows"]]
    assert data["inactive_hidden"] == 3 + 1
    service._CACHE.clear()
    assert "Bengali Chana Sattu" in [r["product"] for r in (await service.build(db, include_inactive=True))["rows"]]


# ── the layout asked for on 12 Oct: "looking too haphazard" ─────────────────────────────────────

def test_category_then_total_stock_then_fba_and_to_purchase_follow_the_product_name():
    from tests.js_harness import run_template_js
    order = run_template_js(ROOT / "templates" / "projections.html", [], ["LEGS", "COLS"],
                            "emit(COLS.map(c => c[0]).slice(0, 6));")
    assert order == ["product", "category", "total_stock", "fba_kg", "to_purchase", "ideal_fba"]


def test_the_lead_times_and_the_settings_start_FOLDED():
    assert 'let leadOpen = localStorage.getItem("projections.leadOpen") === "1";' in PAGE
    assert '<div id="settings-card" hidden' in PAGE
    assert "table.collapsed .lead{display:none}" in PAGE
    # Seasonal and the four legs in COLS, plus the cell markup itself: every editable cell folds.
    assert PAGE.count('"num edit lead"') == 3 and 'class="num edit lead"' in PAGE and '"num lead"' in PAGE


def test_the_formula_commentary_is_gone_and_numbers_are_whole():
    assert "Ideal FBA = projection" not in PAGE
    assert 'function whole(v){ return Math.round(Number(v)).toLocaleString("en-IN"); }' in PAGE
    assert "maximumFractionDigits: 1" not in PAGE


def test_the_profit_tab_shows_whole_percentages_and_kilograms():
    from tests.js_harness import run_portfolio_js
    assert run_portfolio_js('emit([pct(0.3249), kg(1486.5), perKg(100.41)]);') == [
        "32%", "1,487 kg", '<span class="">₹100</span>']


# ── FBA stock and To purchase (13 Oct) ──────────────────────────────────────────────────────────

def test_to_purchase_is_clamped_and_a_dash_when_unknown():
    assert logic.to_purchase(3205.0, 1426.0) == pytest.approx(1779.0)
    assert logic.to_purchase(500.0, 900.0) == 0.0
    assert logic.to_purchase(None, 900.0) is None and logic.to_purchase(500.0, None) is None


def test_the_total_to_purchase_is_the_sum_of_CLAMPED_rows():
    """One product in surplus must never cancel two short ones (the Orders tab's to_buy rule)."""
    rows = [{"to_purchase": 0.0, "fba_kg": 900.0, "total_stock": 100.0},
            {"to_purchase": 100.0, "fba_kg": 0.0, "total_stock": 100.0},
            {"to_purchase": 50.0, "fba_kg": 10.0, "total_stock": 60.0}]
    total = logic.sum_rows(rows)
    assert total["to_purchase"] == 150.0
    assert total["total_stock"] - total["fba_kg"] < 0          # the wrong formula would say "buy nothing"


async def _value(v):
    return v


@pytest.fixture
def stock(monkeypatch):
    """FBA units per child ASIN, the MRP weights, and the child -> row mapping."""
    state = {"units": {"B0CHANA1": 1000, "B0CHANA2": 200, "B0RICE1": 50, "B0ORPHAN": 30, "B0JAU1": 0},
             "weights": {"B0CHANA1": 0.5, "B0CHANA2": 2.0, "B0RICE1": 1.0, "B0ORPHAN": 1.0, "B0JAU1": 1.0},
             "names": {"B0ORPHAN": "Old Thing"},
             "parents": {"B0CHANA1": "B0CHANA", "B0CHANA2": "B0CHANA", "B0RICE1": "B0RICE",
                         "B0ORPHAN": "B0GONE", "B0JAU1": "B0JAU"},
             "run": {"as_of": "2026-10-13T02:15:00", "status": "done", "error": None}}

    async def fake_stock(db):
        return state["units"], state["run"]

    async def fake_catalogue():
        return ({a: {"weight": w, "name": state["names"].get(a, a)} for a, w in state["weights"].items()},
                None, "sheet")

    async def fake_parents(db):
        return state["parents"]

    monkeypatch.setattr(service, "_fba_units", fake_stock)
    monkeypatch.setattr("app.shipment.catalogue.load_catalogue", fake_catalogue)
    monkeypatch.setattr("app.portfolio.repository.parent_of_children", fake_parents)
    monkeypatch.setattr("app.portfolio.repository.flavour_keys", lambda db, cat: _value({}))
    return state


async def test_fba_kg_weighs_each_child_and_to_purchase_follows(profit, stock, db):
    rows = {r["product"]: r for r in (await service.build(db))["rows"]}
    chana = rows["Chana Sattu"]
    assert chana["fba_kg"] == pytest.approx(1000 * 0.5 + 200 * 2.0)            # 900 kg
    assert chana["to_purchase"] == pytest.approx(max(0.0, chana["total_stock"] - 900.0), abs=0.1)
    assert rows["Jau Sattu"]["fba_kg"] == 0.0           # holds nothing: a real zero once a run exists
    assert rows["Brand New Thing"]["fba_kg"] == 0.0      # not in the run at all: also nothing held


async def test_stock_for_a_product_with_no_row_is_named_not_dropped(profit, stock, db):
    data = await service.build(db)
    assert data["fba_stock"]["unmapped"] == [{"product": "Old Thing", "kg": 30.0}]
    assert data["fba_stock"]["as_of"] == "2026-10-13T02:15:00" and data["fba_stock"]["status"] == "done"


async def test_no_good_run_is_a_dash_never_buy_everything(profit, stock, db):
    stock["units"], stock["run"] = None, {"as_of": None, "status": None, "error": None}
    rows = (await service.build(db))["rows"]
    assert rows and all(r["fba_kg"] is None and r["to_purchase"] is None for r in rows)


async def test_an_unknown_pack_weight_makes_that_rows_fba_kg_a_dash(profit, stock, db):
    stock["weights"]["B0CHANA2"] = None
    data = await service.build(db)
    chana = next(r for r in data["rows"] if r["product"] == "Chana Sattu")
    assert chana["fba_kg"] is None and chana["to_purchase"] is None
    assert data["fba_stock"]["unknown_weight"] == ["B0CHANA2"]


async def test_a_real_stock_run_feeds_the_table(profit, db, monkeypatch):
    """End to end through the stored run, not a stand-in for `_fba_units`."""
    from app.inventory import refresh

    async def get(path, params):
        return {"payload": {"inventorySummaries": [
            {"sellerSku": "1kg cs FBA", "asin": "B0CHANA1",
             "inventoryDetails": {"fulfillableQuantity": 10, "inboundShippedQuantity": 30}},
            {"sellerSku": "1kg cs flex", "asin": "B0CHANA1", "inventoryDetails": {"fulfillableQuantity": 999}}]}}
    await refresh.run("button", get=get)

    async def cat():
        return {"B0CHANA1": {"weight": 1.0, "name": "Chana Sattu"}}, None, "sheet"
    monkeypatch.setattr("app.shipment.catalogue.load_catalogue", cat)
    monkeypatch.setattr("app.portfolio.repository.parent_of_children", lambda db: _value({"B0CHANA1": "B0CHANA"}))
    monkeypatch.setattr("app.portfolio.repository.flavour_keys", lambda db, c: _value({}))
    chana = next(r for r in (await service.build(db))["rows"] if r["product"] == "Chana Sattu")
    assert chana["fba_kg"] == 40.0


async def test_refresh_stock_reads_amazon_and_returns_the_new_table(profit, auth_client, monkeypatch):
    calls = []

    async def fake_run(source, **kw):
        calls.append(source)
        return {"id": 1, "status": "done", "sku_count": 1, "units": 5, "error": None}
    monkeypatch.setattr("app.inventory.refresh.run", fake_run)
    r = await auth_client.post("/projections/refresh-stock")
    assert r.status_code == 200 and calls == ["button"] and "rows" in r.json()


async def test_refresh_stock_while_one_runs_is_409(profit, auth_client, monkeypatch):
    async def busy(source, **kw):
        return {"refused": True}
    monkeypatch.setattr("app.inventory.refresh.run", busy)
    assert (await auth_client.post("/projections/refresh-stock")).status_code == 409


async def test_refresh_stock_needs_the_projections_area(ops_client):
    assert (await ops_client.post("/projections/refresh-stock")).status_code in (303, 403)


def test_a_missing_or_failed_stock_read_gets_a_banner_and_a_good_one_does_not():
    from tests.js_harness import run_template_js
    out = run_template_js(ROOT / "templates" / "projections.html", ["renderMessages", "stockTime", "esc", "kg", "whole"], [], """
const box = {innerHTML: ""};
$ = () => box;
const states = [{}, {as_of: "2026-10-13T02:15:00", status: "failed", error: "Amazon 503"},
                {as_of: "2026-10-13T02:15:00", status: "done", stale: false},
                {as_of: "2026-10-13T02:15:00", status: "done", stale: false, unmapped: [{product: "Old Thing", kg: 30}]}];
emit(states.map(fs => { data = {complete_30: true, complete_7: true, fba_stock: fs}; renderMessages(); return box.innerHTML; }));
""")
    assert "not been read" in out[0] and "Amazon 503" in out[1] and out[2] == "" and "Old Thing 30 kg" in out[3]


async def test_no_stock_is_named_as_unmapped_while_the_table_itself_is_missing(profit, stock, db, monkeypatch):
    from app.routers import portfolio as pf
    inner = pf._dashboard

    async def incomplete(db, window, *, include_inactive=False):
        out = await inner(db, window, include_inactive=include_inactive)
        return {**out, "completeness": {"complete": False, "missing_count": 30}}
    monkeypatch.setattr("app.routers.portfolio._dashboard", incomplete)
    data = await service.build(db)
    assert data["rows"] == [] and data["fba_stock"]["unmapped"] == []
