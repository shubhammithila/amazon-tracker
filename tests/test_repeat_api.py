from datetime import timedelta

import pytest

from app import ist
from app.repeat import repository

pytestmark = pytest.mark.regression


def _line(order, buyer, day, child="B0CHILD001", parent="B0PARENT01"):
    return {"amazon_order_id": order, "shipment_item_id": "1", "buyer_key": buyer,
            "purchase_day": day.isoformat(), "seller_sku": "s", "child_asin": child,
            "parent_asin": parent, "units": 1}


async def _seed(db, buyers=25, back=5, howrah=False):
    as_of = ist.yesterday() - timedelta(days=3)
    first = as_of - timedelta(days=59)                     # inside the 30-day window's period
    lines = []
    for i in range(buyers):
        lines.append(_line(f"a{i}", f"k{i}", first))
        if i < back:                                       # come back 10 days later
            lines.append(_line(f"b{i}", f"k{i}", first + timedelta(days=10)))
    if howrah:
        lines += [_line(f"h{i}", f"h{i}", first, child="B0HOWRAH01", parent="B0HOWRAHP1")
                  for i in range(30)]
    await repository.save_lines(db, lines)
    await repository.record_run(db, window_start=(as_of - timedelta(days=400)).isoformat(),
                                window_end=(as_of + timedelta(days=3)).isoformat(), status="done")


@pytest.fixture
def catalogue(monkeypatch):
    cat = {"B0CHILD001": {"name": "Chana Sattu", "brand": "Mithila Foods", "fba_sku": "s"},
           "B0HOWRAH01": {"name": "Bengali Posta", "brand": "Howrah Foods"}}

    async def fake():
        return cat, None, "sheet"
    monkeypatch.setattr("app.repeat.service.load_catalogue", fake)
    return cat


async def test_the_route_returns_the_30_day_same_repeat_from_counts(auth_client, db, catalogue):
    await _seed(db)
    data = (await auth_client.get("/portfolio/repeat")).json()
    row = next(r for r in data["rows"] if r["parent_asin"] == "B0PARENT01")
    assert row["product"] == "Chana Sattu" and row["w"]["30"]["buyers"] == 25
    assert row["w"]["30"]["same_pct"] == pytest.approx(5 / 25)
    # 25 first orders + 5 returns, one unit each: the 5 repeaters bought 10 of the 30 units.
    assert row["w"]["30"]["units_pct"] == pytest.approx(10 / 30)
    assert data["total"]["30"] == {"buyers": 25, "repeat_pct": pytest.approx(5 / 25),
                                   "units_pct": pytest.approx(10 / 30)}
    assert data["brand"] == "Mithila Foods"


async def test_a_small_cohort_returns_None_not_zero(auth_client, db, catalogue):
    await _seed(db, buyers=10)
    row = (await auth_client.get("/portfolio/repeat")).json()["rows"][0]
    assert row["w"]["30"]["buyers"] == 10 and row["w"]["30"]["same_pct"] is None


async def test_the_brand_filter_keeps_other_brands_out(auth_client, db, catalogue):
    await _seed(db, howrah=True)
    mithila = (await auth_client.get("/portfolio/repeat")).json()
    assert [r["parent_asin"] for r in mithila["rows"]] == ["B0PARENT01"]
    assert mithila["brands"] == ["Howrah Foods", "Mithila Foods"]
    howrah = (await auth_client.get("/portfolio/repeat?brand=Howrah%20Foods")).json()
    assert [r["product"] for r in howrah["rows"]] == ["Bengali Posta"]
    assert howrah["total"]["30"]["buyers"] == 30


async def test_nothing_stored_yet_says_so_instead_of_zeros(auth_client, catalogue):
    data = (await auth_client.get("/portfolio/repeat")).json()
    assert data["rows"] == [] and data["as_of"] is None


async def test_as_of_leaves_the_newest_days_for_late_shipments(auth_client, db, catalogue):
    await _seed(db)
    data = (await auth_client.get("/portfolio/repeat")).json()
    assert data["as_of"] == (ist.yesterday() - timedelta(days=3)).isoformat()


async def test_the_payload_never_carries_a_buyer_key(auth_client, db, catalogue):
    await _seed(db)
    text = (await auth_client.get("/portfolio/repeat")).text
    assert '"k1"' not in text and "buyer_key" not in text


async def test_the_route_needs_a_login(client):
    r = await client.get("/portfolio/repeat", follow_redirects=False)
    assert r.status_code in (302, 303, 401)


def _econ(child, parent, sales):
    return {"parentAsin": parent, "childAsin": child,
            "sales": {"orderedProductSales": {"amount": sales}, "unitsOrdered": 10,
                      "netUnitsSold": 10, "refundedProductSales": {"amount": 0}},
            "fees": [], "ads": [], "netProceeds": {"total": {"amount": sales / 2}}}


def test_names_match_the_profit_view_for_the_same_children():
    """The same product must read the same on both sub-tabs, including a derived family name and
    the collision suffix the Profit view adds."""
    from app.portfolio import logic as pf
    from app.repeat import logic
    catalogue = {
        "B0RC000001": {"name": "Peri Peri Roasted Chana", "brand": "Mithila Foods", "weight": 0.5},
        "B0RC000002": {"name": "Nimbu Pudina Roasted Chana", "brand": "Mithila Foods", "weight": 0.5},
        "B0RC000003": {"name": "Roasted Chana", "brand": "Mithila Foods", "weight": 1.0},
        "B0CS000001": {"name": "Chana Sattu", "brand": "Mithila Foods", "weight": 1.0},
        "B0CS000002": {"name": "Chana Sattu", "brand": "Mithila Foods", "weight": 0.5},
    }
    rows = [_econ("B0RC000001", "B0RCPAR001", 900), _econ("B0RC000002", "B0RCPAR001", 300),
            _econ("B0RC000003", "B0RCPAR002", 500), _econ("B0CS000001", "B0CSPAR001", 800),
            _econ("B0CS000002", "B0CSPAR001", 200)]
    data = pf.portfolio(rows, catalogue, {})
    profit = {p["parent_asin"]: p["product"] for p in data["parents"]}
    children = {p["parent_asin"]: [catalogue[s["asin"]]["name"] for s in p["sizes"]]
                for p in data["parents"]}
    assert logic.name_parents(children) == profit
    assert profit["B0RCPAR001"] == "Roasted Chana (2 flavours)"


async def test_a_fetch_that_includes_TODAY_still_ends_as_of_at_yesterday_minus_the_lag(
        auth_client, db, catalogue):
    await _seed(db)
    await repository.record_run(db, window_start=ist.today().isoformat(),
                                window_end=ist.today().isoformat(), status="done")
    data = (await auth_client.get("/portfolio/repeat")).json()
    assert data["as_of"] == (ist.yesterday() - timedelta(days=3)).isoformat()


# ── categories and the download ────────────────────────────────────────────────────────────────

async def _classify(db, name="chana sattu", priority=1):
    from app.models import ProductCategory
    db.add(ProductCategory(product_key=name, priority=priority))
    await db.commit()


async def test_rows_carry_the_profit_views_category_and_each_category_has_its_own_total(
        auth_client, db, catalogue):
    await _seed(db, howrah=True)
    await _classify(db)
    data = (await auth_client.get("/portfolio/repeat")).json()
    row = next(r for r in data["rows"] if r["parent_asin"] == "B0PARENT01")
    assert row["category"] == "Sattu"
    sattu = next(c for c in data["categories"] if c["category"] == "Sattu")
    assert sattu["products"] == 1 and sattu["total"]["30"]["buyers"] == 25
    assert sattu["total"]["30"]["repeat_pct"] == pytest.approx(5 / 25)
    # Howrah's product is a different brand, so it is in neither the rows nor the categories.
    assert all(c["category"] != "Unclassified" for c in data["categories"])


async def test_an_unclassified_product_is_its_own_bucket_never_guessed(auth_client, db, catalogue):
    await _seed(db)
    data = (await auth_client.get("/portfolio/repeat")).json()
    assert data["rows"][0]["category"] == "Unclassified"
    assert [c["category"] for c in data["categories"]] == ["Unclassified"]


async def test_the_excel_holds_the_screens_rows_as_numbers_90_days_first(auth_client, db, catalogue):
    import io

    from openpyxl import load_workbook
    await _seed(db)
    r = await auth_client.post("/portfolio/repeat/export",
                               json={"format": "xlsx", "ids": ["B0PARENT01"]})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    heads = [c.value for c in ws[2]]
    assert heads[:3] == ["Product", "Category", "FBA share"]
    assert heads.index("90d buyers") < heads.index("60d buyers") < heads.index("30d buyers")
    assert ws.freeze_panes == "B3"
    assert ws.cell(1, 1).value == "Mithila Foods — all products"
    assert ws.cell(3, 1).value == "Chana Sattu"
    same30 = ws.cell(3, heads.index("30d same repeat") + 1)
    assert same30.value == pytest.approx(5 / 25) and same30.number_format == "0.0%"
    units30 = ws.cell(3, heads.index("30d repeat units") + 1)
    assert units30.value == pytest.approx(10 / 30)


async def test_a_category_download_is_totalled_for_that_category(auth_client, db, catalogue):
    import io

    from openpyxl import load_workbook
    await _seed(db)
    await _classify(db)
    r = await auth_client.post("/portfolio/repeat/export",
                               json={"format": "xlsx", "ids": ["B0PARENT01"], "category": "Sattu"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(1, 1).value == "Sattu — all products"
    assert "sattu" in r.headers["content-disposition"]


async def test_the_pdf_downloads_and_bad_requests_are_refused(auth_client, db, catalogue):
    await _seed(db)
    r = await auth_client.post("/portfolio/repeat/export", json={"format": "pdf", "ids": ["B0PARENT01"]})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert (await auth_client.post("/portfolio/repeat/export", json={"format": "csv"})).status_code == 400
    assert (await auth_client.post("/portfolio/repeat/export",
                                   json={"format": "xlsx", "ids": [1]})).status_code == 400


def test_an_unavailable_window_downloads_as_blanks_not_its_hidden_figures():
    from app.repeat import export
    payload = {"windows": {"90": {"available": False}, "60": {"available": True},
                           "30": {"available": True}},
               "rows": [{"parent_asin": "P", "product": "X", "category": "Rest", "fba_share": 1.0,
                         "w": {"90": {"buyers": 99, "same_pct": 0.5}, "60": {"buyers": 40, "same_pct": 0.1}}}],
               "total": {}, "categories": [], "brand": "B"}
    values = export.build_table(payload, ["P"], None).rows[0].values
    assert values["buyers-90"] is None and values["same-90"] is None
    assert values["buyers-60"] == 40



def test_the_download_keeps_the_screens_order_and_only_its_rows():
    """The screen sends the ids it shows, sorted and filtered; the file must be those rows exactly."""
    from app.repeat import export
    row = lambda pid, name: {"parent_asin": pid, "product": name, "category": "Rest",  # noqa: E731
                             "fba_share": 1.0, "w": {}}
    payload = {"windows": {}, "rows": [row("A", "Alpha"), row("B", "Beta"), row("C", "Gamma")],
               "total": {}, "categories": [], "brand": "B"}
    got = [r.values["product"] for r in export.build_table(payload, ["C", "A"], None).rows]
    assert got == ["Gamma", "Alpha"]
