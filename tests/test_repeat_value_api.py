"""Customer value through the real routes: payload, page, downloads."""
import io
from datetime import timedelta

import pytest

from app import ist
from app.repeat import repository

pytestmark = pytest.mark.regression


@pytest.fixture
def catalogue(monkeypatch):
    cat = {"B0CHILD001": {"name": "Chana Sattu", "brand": "Mithila Foods"}}

    async def fake():
        return cat, None, "sheet"
    monkeypatch.setattr("app.repeat.service.load_catalogue", fake)
    return cat


async def _seed(db, n=25):
    as_of = ist.yesterday() - timedelta(days=3)
    first = as_of - timedelta(days=150)
    lines = []
    for i in range(n):
        for k, (off, rev) in enumerate(((0, 200.0), (20, 150.0))):
            lines.append({"amazon_order_id": f"o{i}-{k}", "shipment_item_id": "1",
                          "buyer_key": f"k{i}", "purchase_day": (first + timedelta(days=off)).isoformat(),
                          "seller_sku": "s", "child_asin": "B0CHILD001", "parent_asin": "B0PARENT01",
                          "units": 1, "revenue": rev})
    await repository.save_lines(db, lines)
    await repository.record_run(db, window_start="2026-01-01",
                                window_end=(as_of + timedelta(days=3)).isoformat(), status="done")


async def test_the_payload_carries_ltv_per_first_product_and_the_brand_total(auth_client, db, catalogue):
    await _seed(db)
    d = (await auth_client.get("/portfolio/value")).json()
    row = d["rows"][0]
    assert row["product"] == "Chana Sattu" and row["new_customers"] == 25
    assert row["ltv"]["30"] == 350.0 and d["total"]["ltv"]["30"] == 350.0
    assert d["brands"][0] == "All brands" and d["horizons"] == [30, 60, 90, 180]
    assert d["categories"][0]["category"] == "Unclassified"


async def test_unpriced_lines_hold_ltv_back_and_the_screen_says_from_when(auth_client, db, catalogue):
    await _seed(db)
    from sqlalchemy import update
    from app.models import CustomerOrderLine
    await db.execute(update(CustomerOrderLine).values(revenue=None))
    await db.commit()
    d = (await auth_client.get("/portfolio/value")).json()
    assert d["priced_from"] is None and d["rows"][0]["ltv"]["30"] is None


async def test_the_page_is_the_third_portfolio_sub_tab(auth_client):
    html = (await auth_client.get("/portfolio-page/value")).text
    assert 'href="/portfolio-page/value"' in html and "Customer value" in html
    for page in ("/portfolio-page", "/portfolio-page/repeat"):
        assert 'href="/portfolio-page/value"' in (await auth_client.get(page)).text


async def test_the_downloads_hold_the_rows_as_numbers(auth_client, db, catalogue):
    from openpyxl import load_workbook
    await _seed(db)
    r = await auth_client.post("/portfolio/value/export", json={"format": "xlsx", "ids": ["B0PARENT01"]})
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    heads = [c.value for c in ws[2]]
    assert heads[:4] == ["Product", "Category", "New customers", "LTV 30d"]
    assert ws.cell(3, 1).value == "Chana Sattu" and ws.cell(3, 4).value == 350.0
    assert ws.cell(1, 1).value == "Mithila Foods — all products"
    pdf = await auth_client.post("/portfolio/value/export", json={"format": "pdf", "ids": ["B0PARENT01"]})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert (await auth_client.post("/portfolio/value/export", json={"format": "csv"})).status_code == 400



async def test_the_payload_is_cached_until_the_data_moves(auth_client, db, catalogue):
    from app.repeat import value_service
    await _seed(db)
    first = (await auth_client.get("/portfolio/value")).json()
    calls = []
    real = value_service._build

    async def counting(*a, **k):
        calls.append(1)
        return await real(*a, **k)
    value_service._build = counting
    try:
        assert (await auth_client.get("/portfolio/value")).json() == first and calls == []
        await repository.record_run(db, window_start="2026-10-01", window_end="2026-10-02", status="done")
        await auth_client.get("/portfolio/value")
        assert calls == [1], "new data must rebuild the payload"
    finally:
        value_service._build = real
